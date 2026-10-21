from __future__ import annotations

import csv, json, os, re, statistics, threading, time, uuid
from dataclasses import dataclass, field
from enum import Enum
from typing import Any, Dict, List, Optional


class QueryType(str, Enum):
    FILTER = "filter"
    AGGREGATE = "aggregate"
    SORT = "sort"
    SELECT = "select"
    COMPARE = "compare"
    SUMMARIZE = "summarize"
    SEARCH = "search"
    CALCULATE = "calculate"
    GROUP = "group"


@dataclass
class ParsedQuery:
    query_id: str = field(default_factory=lambda: str(uuid.uuid4()))
    raw_text: str = ""
    query_type: QueryType = QueryType.SELECT
    target_table: Optional[str] = None
    columns: List[str] = field(default_factory=list)
    conditions: List[Dict[str, Any]] = field(default_factory=list)
    aggregations: List[Dict[str, Any]] = field(default_factory=list)
    sort_by: Optional[str] = None
    sort_order: str = "asc"
    limit: Optional[int] = None
    confidence: float = 0.0
    sql_equivalent: Optional[str] = None
    explanation: str = ""

    def to_dict(self) -> Dict[str, Any]:
        return {
            "query_id": self.query_id, "raw_text": self.raw_text,
            "query_type": self.query_type.value, "target_table": self.target_table,
            "columns": self.columns, "conditions": self.conditions,
            "aggregations": self.aggregations, "sort_by": self.sort_by,
            "sort_order": self.sort_order, "limit": self.limit,
            "confidence": self.confidence, "sql_equivalent": self.sql_equivalent,
            "explanation": self.explanation,
        }


@dataclass
class QueryResult:
    result_id: str = field(default_factory=lambda: str(uuid.uuid4()))
    query_id: str = ""
    success: bool = False
    data: List[Dict[str, Any]] = field(default_factory=list)
    row_count: int = 0
    column_names: List[str] = field(default_factory=list)
    execution_time_ms: float = 0.0
    error: Optional[str] = None
    summary: Optional[str] = None

    def to_dict(self) -> Dict[str, Any]:
        return {
            "result_id": self.result_id, "query_id": self.query_id,
            "success": self.success, "data": self.data, "row_count": self.row_count,
            "column_names": self.column_names, "execution_time_ms": self.execution_time_ms,
            "error": self.error, "summary": self.summary,
        }


@dataclass
class DataSummary:
    summary_id: str = field(default_factory=lambda: str(uuid.uuid4()))
    file_path: str = ""
    total_rows: int = 0
    total_columns: int = 0
    columns: List[Dict[str, Any]] = field(default_factory=list)
    statistics: Dict[str, Any] = field(default_factory=dict)
    key_insights: List[str] = field(default_factory=list)

    def to_dict(self) -> Dict[str, Any]:
        return {
            "summary_id": self.summary_id, "file_path": self.file_path,
            "total_rows": self.total_rows, "total_columns": self.total_columns,
            "columns": self.columns, "statistics": self.statistics,
            "key_insights": self.key_insights,
        }


class NLQEngine:
    VERSION = "1.0.0"
    _AGG = {"sum": "SUM", "total": "SUM", "count": "COUNT", "number of": "COUNT",
            "how many": "COUNT", "average": "AVG", "avg": "AVG", "mean": "AVG",
            "minimum": "MIN", "min": "MIN", "lowest": "MIN",
            "maximum": "MAX", "max": "MAX", "highest": "MAX", "largest": "MAX"}
    _CMP = {"greater than": ">", "more than": ">", "above": ">",
            "less than": "<", "fewer than": "<", "below": "<", "under": "<",
            "equals": "=", "equal to": "=", "is": "=", "was": "=",
            "not equal": "!=", "different from": "!=",
            "at least": ">=", "at most": "<=", "contains": "LIKE", "like": "LIKE"}

    def __init__(self) -> None:
        self._query_history: List[Dict[str, Any]] = []
        self._data_sources: Dict[str, str] = {}
        self._parsed_schemas: Dict[str, Dict[str, Any]] = {}
        self._lock = threading.RLock()

    def register_data_source(self, name: str, file_path: str) -> None:
        with self._lock:
            if os.path.exists(file_path):
                self._data_sources[name] = file_path
                self._parsed_schemas[name] = self._schema(file_path)

    def _schema(self, path: str) -> Dict[str, Any]:
        ext = os.path.splitext(path)[1].lower()
        try:
            if ext == ".csv":
                with open(path, "r", encoding="utf-8", errors="replace") as f:
                    rows = list(csv.DictReader(f))
                    return {"columns": list(rows[0].keys()) if rows else [], "row_count": len(rows)}
            elif ext == ".json":
                with open(path, "r", encoding="utf-8") as f:
                    data = json.load(f)
                    if isinstance(data, list) and data:
                        return {"columns": list(data[0].keys()), "row_count": len(data)}
            elif ext in (".xlsx", ".xls"):
                import openpyxl
                wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
                try:
                    ws = wb.active
                    rows = list(ws.iter_rows(values_only=True))
                finally:
                    # Item 25A (C34): close on every path (best-effort);
                    # the old close ran only after a successful read.
                    try:
                        wb.close()
                    except Exception:
                        pass
                if rows:
                    hdrs = [str(h) if h else f"col_{i}" for i, h in enumerate(rows[0])]
                    return {"columns": hdrs, "row_count": len(rows) - 1}
        except Exception:
            pass
        return {"columns": [], "row_count": 0}

    def _load(self, path: str) -> List[Dict[str, Any]]:
        ext = os.path.splitext(path)[1].lower()
        try:
            if ext == ".csv":
                with open(path, "r", encoding="utf-8", errors="replace") as f:
                    return list(csv.DictReader(f))
            elif ext == ".json":
                with open(path, "r", encoding="utf-8") as f:
                    d = json.load(f)
                    return d if isinstance(d, list) else []
            elif ext in (".xlsx", ".xls"):
                import openpyxl
                wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
                try:
                    ws = wb.active
                    rows = list(ws.iter_rows(values_only=True))
                finally:
                    # Item 25A (C34): close on every path (best-effort);
                    # the old close ran only after a successful read.
                    try:
                        wb.close()
                    except Exception:
                        pass
                if len(rows) < 2:
                    return []
                hdrs = [str(h) if h else f"col_{i}" for i, h in enumerate(rows[0])]
                return [dict(zip(hdrs, r)) for r in rows[1:]]
        except Exception:
            pass
        return []

    def _resolve(self, text: str) -> Optional[str]:
        low = text.lower()
        for n in self._data_sources:
            if n.lower() in low:
                return n
        return list(self._data_sources.keys())[0] if self._data_sources else None

    def parse_query(self, text: str) -> ParsedQuery:
        with self._lock:
            low = text.lower().strip()
            qt = self._detect_type(low)
            tbl = self._resolve(text)
            cols = self._get_cols(low, tbl)
            conds = self._get_conds(low)
            aggs = self._get_aggs(low)
            sb = self._get_sort(low)
            so = "desc" if any(w in low for w in ("descending", "desc", "highest", "top")) else "asc"
            lim = self._get_limit(low)
            conf = self._confidence(low, qt, tbl)
            pq = ParsedQuery(raw_text=text, query_type=qt, target_table=tbl, columns=cols,
                             conditions=conds, aggregations=aggs, sort_by=sb, sort_order=so,
                             limit=lim, confidence=conf, explanation=self._explain(qt, cols, conds, aggs))
            pq.sql_equivalent = self._gen_sql(pq)
            return pq

    def _detect_type(self, t: str) -> QueryType:
        if any(k in t for k in self._AGG):
            return QueryType.AGGREGATE
        for w, qt in [("group", QueryType.GROUP), ("sort", QueryType.SORT), ("order by", QueryType.SORT),
                       ("compare", QueryType.COMPARE), ("versus", QueryType.COMPARE),
                       ("summarize", QueryType.SUMMARIZE), ("summary", QueryType.SUMMARIZE),
                       ("search", QueryType.SEARCH), ("find", QueryType.SEARCH),
                       ("calculate", QueryType.CALCULATE), ("compute", QueryType.CALCULATE),
                       ("filter", QueryType.FILTER), ("where", QueryType.FILTER),
                       ("show me", QueryType.FILTER), ("get", QueryType.FILTER), ("list", QueryType.FILTER)]:
            if w in t:
                return qt
        return QueryType.SELECT

    def _get_cols(self, t: str, tbl: Optional[str]) -> List[str]:
        for pat in [r"show\s+(?:me\s+)?(.+?)(?:\s+where|\s+with|\s+group|\s+sort|$)",
                    r"select\s+(.+?)(?:\s+from|\s+where|$)", r"get\s+(.+?)(?:\s+where|\s+from|$)"]:
            m = re.search(pat, t, re.IGNORECASE)
            if m:
                return [p.strip().strip("\"'") for p in re.split(r",\s*|\s+and\s+", m.group(1).strip())
                        if p.strip() and len(p.strip()) < 50][:20]
        if tbl and tbl in self._parsed_schemas:
            return list(self._parsed_schemas[tbl]["columns"])
        return []

    def _get_conds(self, t: str) -> List[Dict[str, Any]]:
        conds = []
        for ph, op in self._CMP.items():
            for m in re.finditer(rf"(\w+)\s+{re.escape(ph)}\s+(['\"]?[^'\"]+['\"]?)", t):
                fld, raw = m.group(1).strip(), m.group(2).strip().strip("\"'")
                try:
                    v: Any = float(raw)
                    if v == int(v): v = int(v)
                except ValueError:
                    v = raw
                conds.append({"field": fld, "operator": op, "value": v})
        if not conds:
            wm = re.search(r"where\s+(.+?)(?:\s+and\s+|\s+or\s+|\s+order|\s+sort|$)", t)
            if wm:
                for ph, op in self._CMP.items():
                    if ph in wm.group(1):
                        parts = wm.group(1).split(ph, 1)
                        if len(parts) == 2:
                            fld = parts[0].strip().split()[-1] if parts[0].strip() else "val"
                            raw = parts[1].strip().split()[0].strip().strip("\"'")
                            try:
                                v = float(raw)
                                if v == int(v): v = int(v)
                            except ValueError:
                                v = raw
                            conds.append({"field": fld, "operator": op, "value": v})
                            break
        return conds

    def _get_aggs(self, t: str) -> List[Dict[str, Any]]:
        aggs = []
        for kw, func in self._AGG.items():
            if kw in t:
                m = re.search(rf"{re.escape(kw)}\s+(?:of\s+)?(\w+)", t)
                col = m.group(1) if m else "*"
                if not any(a["function"] == func and a["column"] == col for a in aggs):
                    aggs.append({"function": func, "column": col})
        return aggs

    def _get_sort(self, t: str) -> Optional[str]:
        for kw in ("sort by", "order by", "rank by"):
            m = re.search(rf"{kw}\s+(\w+)", t)
            if m: return m.group(1)
        return None

    def _get_limit(self, t: str) -> Optional[int]:
        m = re.search(r"(?:top|first|limit|show)\s+(\d+)", t)
        return int(m.group(1)) if m else None

    def _confidence(self, t: str, qt: QueryType, tbl: Optional[str]) -> float:
        s = 0.3
        if tbl and tbl in self._parsed_schemas: s += 0.3
        if qt != QueryType.SELECT: s += 0.15
        if len(t.split()) >= 3: s += 0.1
        if any(op in t for op in self._CMP): s += 0.1
        return min(s, 1.0)

    def _explain(self, qt: QueryType, cols: List[str], conds: List[Dict], aggs: List[Dict]) -> str:
        p = [f"Query type: {qt.value}"]
        if cols: p.append(f"Columns: {', '.join(cols)}")
        if conds: p.append(f"Conditions: {len(conds)} filter(s)")
        if aggs: p.append(f"Aggregations: {', '.join(a['function'] for a in aggs)}")
        return ". ".join(p)

    def _gen_sql(self, q: ParsedQuery) -> str:
        tbl = q.target_table or "data"
        c = ", ".join(q.columns) if q.columns else "*"
        sql = f"SELECT {c} FROM {tbl}"
        if q.conditions:
            wh = [f"{x['field']} {x['operator']} '{x['value']}'" if isinstance(x["value"], str)
                  else f"{x['field']} {x['operator']} {x['value']}" for x in q.conditions]
            sql += " WHERE " + " AND ".join(wh)
        if q.aggregations:
            ap = [f"{a['function']}({a['column']})" for a in q.aggregations]
            sel = ", ".join(q.columns) + ", " + ", ".join(ap) if q.columns else ", ".join(ap)
            sql = f"SELECT {sel} FROM {tbl}"
        if q.sort_by: sql += f" ORDER BY {q.sort_by} {q.sort_order.upper()}"
        if q.limit: sql += f" LIMIT {q.limit}"
        return sql + ";"

    def execute_query(self, query: ParsedQuery) -> QueryResult:
        start = time.time()
        r = QueryResult(query_id=query.query_id)
        with self._lock:
            try:
                if not query.target_table or query.target_table not in self._data_sources:
                    r.error = f"Data source '{query.target_table}' not found"
                    return r
                data = self._load(self._data_sources[query.target_table])
                if query.conditions:
                    data = self._filter(data, query.conditions)
                if query.query_type == QueryType.GROUP and query.columns:
                    data = self._group(data, query.columns)
                if query.aggregations:
                    data = self._agg(data, query.aggregations)
                if query.sort_by:
                    data = self._sort(data, query.sort_by, query.sort_order)
                if query.limit:
                    data = data[:query.limit]
                r.success, r.data, r.row_count = True, data, len(data)
                r.column_names = list(data[0].keys()) if data else []
                self._query_history.append({"query": query.to_dict(), "timestamp": time.time(),
                                            "result_rows": r.row_count})
            except Exception as e:
                r.error = str(e)
            r.execution_time_ms = (time.time() - start) * 1000
            return r

    def _filter(self, data: List[Dict], conds: List[Dict]) -> List[Dict]:
        out = []
        for row in data:
            ok = True
            for c in conds:
                v = row.get(c["field"])
                if v is None: ok = False; break
                e = c["value"]
                try:
                    if isinstance(e, (int, float)): v = float(v)
                    elif isinstance(e, str): v = str(v)
                except (ValueError, TypeError): ok = False; break
                op = c["operator"]
                if op == ">" and not (v > e): ok = False
                elif op == "<" and not (v < e): ok = False
                elif op == ">=" and not (v >= e): ok = False
                elif op == "<=" and not (v <= e): ok = False
                elif op == "=" and v != e: ok = False
                elif op == "!=" and v == e: ok = False
                elif op == "LIKE" and e.lower() not in str(v).lower(): ok = False
                if not ok: break
            if ok: out.append(row)
        return out

    def _agg(self, data: List[Dict], aggs: List[Dict]) -> List[Dict]:
        if not data: return []
        res: Dict[str, Any] = {}
        for a in aggs:
            fn, col = a["function"], a["column"]
            key = f"{fn}({col})"
            vals = []
            for row in data:
                v = row.get(col)
                if v is not None:
                    try: vals.append(float(v))
                    except (ValueError, TypeError):
                        if fn == "COUNT": vals.append(1)
            if fn == "SUM": res[key] = sum(vals) if vals else 0
            elif fn == "COUNT": res[key] = len(vals)
            elif fn == "AVG": res[key] = round(statistics.mean(vals), 4) if vals else 0
            elif fn == "MIN": res[key] = min(vals) if vals else 0
            elif fn == "MAX": res[key] = max(vals) if vals else 0
        return [res] if res else []

    def _sort(self, data: List[Dict], sb: str, order: str) -> List[Dict]:
        def k(r: Dict) -> Any:
            v = r.get(sb, "")
            try: return float(v)
            except (ValueError, TypeError): return str(v).lower()
        return sorted(data, key=k, reverse=(order == "desc"))

    def _group(self, data: List[Dict], cols: List[str]) -> List[Dict]:
        g: Dict[tuple, List[Dict]] = {}
        for row in data:
            k = tuple(str(row.get(c, "")) for c in cols)
            g.setdefault(k, []).append(row)
        return [dict(zip(cols, k)) | {"row_count": len(v)} for k, v in g.items()]

    def summarize_file(self, file_path: str) -> DataSummary:
        with self._lock:
            s = DataSummary(file_path=file_path)
            if not os.path.exists(file_path): return s
            data = self._load(file_path)
            if not data: return s
            s.total_rows, s.total_columns = len(data), len(data[0].keys())
            cols_info, overall = [], {"total_rows": s.total_rows, "total_columns": s.total_columns}
            for col in data[0].keys():
                vals = [r.get(col) for r in data]
                nn = [v for v in vals if v is not None and v != ""]
                ci: Dict[str, Any] = {"name": col, "non_null_count": len(nn),
                    "null_count": len(vals) - len(nn), "unique_count": len(set(str(v) for v in nn))}
                nv = []
                for v in nn:
                    try: nv.append(float(v))
                    except (ValueError, TypeError): pass
                if nv:
                    ci["type"] = "numeric"
                    ci["statistics"] = self._stats(nv)
                else:
                    ci["type"] = "text"
                    if nn:
                        freq: Dict[str, int] = {}
                        for v in nn: freq[str(v)] = freq.get(str(v), 0) + 1
                        top = sorted(freq.items(), key=lambda x: x[1], reverse=True)[:5]
                        ci["top_values"] = [{"value": k, "count": c} for k, c in top]
                cols_info.append(ci)
                overall[col] = ci.get("statistics", {})
            s.columns, s.statistics = cols_info, overall
            s.key_insights = self._insights(cols_info, overall)
            return s

    def _stats(self, vals: List[float]) -> Dict[str, Any]:
        if not vals: return {}
        sv = sorted(vals)
        n = len(sv)
        return {"mean": round(statistics.mean(sv), 4), "median": round(statistics.median(sv), 4),
                "std_dev": round(statistics.stdev(sv), 4) if n > 1 else 0.0,
                "min": sv[0], "max": sv[-1], "count": n}

    def _insights(self, cols: List[Dict], stats: Dict) -> List[str]:
        ins = []
        total = stats.get("total_rows", 0)
        ins.append(f"Dataset: {total} rows, {stats.get('total_columns', 0)} columns")
        nc = [c for c in cols if c.get("type") == "numeric"]
        tc = [c for c in cols if c.get("type") == "text"]
        if nc:
            ins.append(f"{len(nc)} numeric columns")
            for c in nc[:3]:
                st = c.get("statistics", {})
                if st: ins.append(f"'{c['name']}' range: {st.get('min')} to {st.get('max')} (avg: {st.get('mean')})")
        if tc:
            ins.append(f"{len(tc)} text columns")
            for c in tc[:2]:
                top = c.get("top_values", [])
                if top: ins.append(f"'{c['name']}' top value: '{top[0]['value']}' ({top[0]['count']}x)")
        hn = [c["name"] for c in cols if c.get("null_count", 0) > total * 0.1]
        if hn: ins.append(f"High null columns: {', '.join(hn[:3])}")
        return ins

    def suggest_queries(self, file_path: str) -> List[str]:
        with self._lock:
            if not os.path.exists(file_path): return []
            data = self._load(file_path)
            if not data: return []
            cols = list(data[0].keys())
            name = os.path.splitext(os.path.basename(file_path))[0]
            sg = [f"Show me all data from {name}"]
            if cols: sg += [f"Show me {cols[0]} from {name}", f"Sort {name} by {cols[0]}"]
            num = [c for c in cols if self._num_col(data, c)]
            if num:
                nc = num[0]
                sg += [f"What is the average {nc} in {name}?", f"What is the total {nc} in {name}?",
                       f"Show me the top 10 rows by {nc}"]
                if len(num) > 1: sg.append(f"Compare {num[0]} vs {num[1]}")
            if len(cols) >= 2: sg += [f"Group {name} by {cols[0]}", f"Show me {cols[0]} and {cols[1]} from {name}"]
            sg += [f"How many rows are in {name}?", f"Summarize {name}"]
            return sg[:15]

    def _num_col(self, data: List[Dict], col: str) -> bool:
        for row in data[:10]:
            try: float(row.get(col, "")); return True
            except (ValueError, TypeError): return False
        return False

    def get_query_history(self, limit: int = 50) -> List[Dict[str, Any]]:
        with self._lock: return self._query_history[-limit:]

    def get_engine_summary(self) -> Dict[str, Any]:
        with self._lock:
            return {"version": self.VERSION, "data_sources": list(self._data_sources.keys()),
                    "total_queries": len(self._query_history),
                    "schemas_parsed": list(self._parsed_schemas.keys()),
                    "supported_query_types": [qt.value for qt in QueryType],
                    "supported_file_formats": ["csv", "json", "xlsx"]}
