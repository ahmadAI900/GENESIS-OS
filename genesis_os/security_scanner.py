from __future__ import annotations

import os, re, json, time, threading, hashlib, csv, zipfile
from dataclasses import dataclass, field
from enum import Enum
from typing import Any, Dict, List, Optional, Tuple

try:
    from docx import Document as DocxDocument
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

try:
    from pypdf import PdfReader
    HAS_PYPDF = True
except ImportError:
    HAS_PYPDF = False

try:
    from openpyxl import load_workbook
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


class SensitiveCategory(str, Enum):
    CREDENTIAL = "credential"
    PII = "pii"
    FINANCIAL = "financial"
    MEDICAL = "medical"
    CLASSIFICATION = "classification"
    PROPRIETARY = "proprietary"
    CONTACT = "contact"
    LOCATION = "location"


class SensitivityLevel(str, Enum):
    PUBLIC = "public"
    INTERNAL = "internal"
    CONFIDENTIAL = "confidential"
    SECRET = "secret"
    TOP_SECRET = "top_secret"


@dataclass
class SensitiveFinding:
    finding_id: str
    category: SensitiveCategory
    pattern_name: str
    matched_text: str
    file_path: str
    line_number: Optional[int]
    context: str
    confidence: float
    recommendation: str

    def to_dict(self) -> Dict[str, Any]:
        return {
            "finding_id": self.finding_id, "category": self.category.value,
            "pattern_name": self.pattern_name, "matched_text": self.matched_text,
            "file_path": self.file_path, "line_number": self.line_number,
            "context": self.context, "confidence": self.confidence,
            "recommendation": self.recommendation,
        }


@dataclass
class MetadataItem:
    item_id: str
    file_path: str
    metadata_type: str
    key: str
    value: str
    risk: str
    recommendation: str

    def to_dict(self) -> Dict[str, Any]:
        return {
            "item_id": self.item_id, "file_path": self.file_path,
            "metadata_type": self.metadata_type, "key": self.key,
            "value": self.value, "risk": self.risk,
            "recommendation": self.recommendation,
        }


@dataclass
class SecurityReport:
    report_id: str
    file_path: str
    scan_timestamp: float
    overall_risk: SensitivityLevel
    findings: List[SensitiveFinding]
    metadata_items: List[MetadataItem]
    hidden_content_found: bool
    total_findings: int
    risk_score: float
    recommendations: List[str]

    def to_dict(self) -> Dict[str, Any]:
        return {
            "report_id": self.report_id, "file_path": self.file_path,
            "scan_timestamp": self.scan_timestamp,
            "overall_risk": self.overall_risk.value,
            "findings": [f.to_dict() for f in self.findings],
            "metadata_items": [m.to_dict() for m in self.metadata_items],
            "hidden_content_found": self.hidden_content_found,
            "total_findings": self.total_findings,
            "risk_score": self.risk_score,
            "recommendations": self.recommendations,
        }


class SecurityScanner:
    VERSION = "1.0.0"

    def __init__(self) -> None:
        self._scan_history: List[Dict[str, Any]] = []
        self._custom_patterns: Dict[str, Dict[str, Any]] = {}
        self._lock = threading.RLock()
        self._max_history = 500
        self._builtin_patterns: Dict[str, Dict[str, Any]] = {}
        self._register_builtin_patterns()

    def _register_builtin_patterns(self) -> None:
        with self._lock:
            self._builtin_patterns = {
                "ssn": {"pattern": r"\b\d{3}-\d{2}-\d{4}\b", "category": SensitiveCategory.PII,
                        "confidence": 0.9, "recommendation": "Remove or mask SSN in documents"},
                "credit_card_visa": {"pattern": r"\b4\d{3}[\s-]?\d{4}[\s-]?\d{4}[\s-]?\d{4}\b",
                                     "category": SensitiveCategory.FINANCIAL,
                                     "confidence": 0.85, "recommendation": "Remove credit card numbers from documents"},
                "credit_card_mastercard": {"pattern": r"\b5[1-5]\d{2}[\s-]?\d{4}[\s-]?\d{4}[\s-]?\d{4}\b",
                                           "category": SensitiveCategory.FINANCIAL,
                                           "confidence": 0.85, "recommendation": "Remove credit card numbers from documents"},
                "credit_card_amex": {"pattern": r"\b3[47]\d{2}[\s-]?\d{6}[\s-]?\d{5}\b",
                                     "category": SensitiveCategory.FINANCIAL,
                                     "confidence": 0.85, "recommendation": "Remove credit card numbers from documents"},
                "email": {"pattern": r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b",
                          "category": SensitiveCategory.CONTACT,
                          "confidence": 0.95, "recommendation": "Review email addresses for PII exposure"},
                "phone_us": {"pattern": r"\b(\+?1[-.\s]?)?\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}\b",
                             "category": SensitiveCategory.CONTACT,
                             "confidence": 0.8, "recommendation": "Remove phone numbers from public documents"},
                "api_key_generic": {"pattern": r"\b(sk-|ak-|key-)[A-Za-z0-9]{20,}\b",
                                    "category": SensitiveCategory.CREDENTIAL,
                                    "confidence": 0.9, "recommendation": "Remove API keys and rotate immediately"},
                "password_assignment": {"pattern": r"(?i)(password|passwd|pwd|secret|token)\s*[:=]\s*['\"]?[^\s'\"]{4,}['\"]?",
                                        "category": SensitiveCategory.CREDENTIAL,
                                        "confidence": 0.85, "recommendation": "Remove hardcoded passwords from configuration"},
                "aws_access_key": {"pattern": r"\b(AKIA|ABIA|ACCA|ASIA)[A-Z0-9]{16}\b",
                                   "category": SensitiveCategory.CREDENTIAL,
                                   "confidence": 0.95, "recommendation": "Remove AWS access keys and rotate immediately"},
                "private_key": {"pattern": r"-----BEGIN\s+(RSA|DSA|EC|OPENSSH)\s+PRIVATE\s+KEY-----",
                                "category": SensitiveCategory.CREDENTIAL,
                                "confidence": 0.99, "recommendation": "Remove private keys from code and use secure storage"},
                "ip_address": {"pattern": r"\b(?:\d{1,3}\.){3}\d{1,3}\b",
                               "category": SensitiveCategory.LOCATION,
                               "confidence": 0.6, "recommendation": "Review IP addresses for internal infrastructure exposure"},
                "url_with_credentials": {"pattern": r"https?://[^:]+:[^@]+@[^\s]+",
                                         "category": SensitiveCategory.CREDENTIAL,
                                         "confidence": 0.95, "recommendation": "Remove embedded credentials from URLs"},
                "medical_record_number": {"pattern": r"\b(MRN|mrn|Patient\s*ID|patient\s*#)\s*[:#]?\s*\d{6,12}\b",
                                          "category": SensitiveCategory.MEDICAL,
                                          "confidence": 0.8, "recommendation": "Remove medical record numbers to ensure HIPAA compliance"},
                "date_of_birth": {"pattern": r"(?i)(DOB|Date\s*of\s*Birth|Born)\s*[:=]?\s*\d{1,2}[/\-\.]\d{1,2}[/\-\.]\d{2,4}",
                                  "category": SensitiveCategory.PII,
                                  "confidence": 0.85, "recommendation": "Remove dates of birth to prevent identity exposure"},
            }

    def _generate_id(self) -> str:
        return hashlib.md5(f"{time.time_ns()}{os.urandom(8).hex()}".encode()).hexdigest()[:16]

    def _mask_sensitive_text(self, matched: str) -> str:
        if len(matched) <= 4:
            return "*" * len(matched)
        return "*" * (len(matched) - 4) + matched[-4:]

    def _extract_context(self, text: str, start: int, end: int, window: int = 40) -> str:
        s, e = max(0, start - window), min(len(text), end + window)
        ctx = text[s:e].replace("\n", " ").replace("\r", "")
        return f"{'...' if s > 0 else ''}{ctx}{'...' if e < len(text) else ''}"

    def _scan_content(self, file_path: str, content: str) -> List[SensitiveFinding]:
        findings: List[SensitiveFinding] = []
        all_patterns = {**self._builtin_patterns, **self._custom_patterns}
        lines = content.split("\n")
        for name, cfg in all_patterns.items():
            try:
                compiled = re.compile(cfg["pattern"], re.IGNORECASE)
            except re.error:
                continue
            for line_idx, line in enumerate(lines, 1):
                for m in compiled.finditer(line):
                    findings.append(SensitiveFinding(
                        finding_id=self._generate_id(), category=cfg["category"],
                        pattern_name=name, matched_text=self._mask_sensitive_text(m.group()),
                        file_path=file_path, line_number=line_idx,
                        context=self._extract_context(line, m.start(), m.end()),
                        confidence=cfg.get("confidence", 0.5),
                        recommendation=cfg.get("recommendation", "Review and remediate"),
                    ))
        return findings

    def _scan_metadata(self, file_path: str) -> List[MetadataItem]:
        items: List[MetadataItem] = []
        ext = os.path.splitext(file_path)[1].lower()
        if ext == ".docx" and HAS_DOCX:
            try:
                props = DocxDocument(file_path).core_properties
                for fname, key, val, risk in [
                    ("author", "author", props.author, "high"),
                    ("created_date", "created_date", str(props.created), "medium"),
                    ("modified_date", "modified_date", str(props.modified), "medium"),
                ]:
                    if val:
                        items.append(MetadataItem(self._generate_id(), file_path, fname, key, str(val), risk,
                                                  "Consider removing author information before sharing"))
            except Exception:
                pass
        elif ext == ".pdf" and HAS_PYPDF:
            try:
                info = PdfReader(file_path).metadata or {}
                for fname, key in [("author", "author"), ("creator", "creator"), ("created_date", "/CreationDate"),
                                   ("modified_date", "/ModDate"), ("subject", "/Subject"), ("title", "/Title")]:
                    val = str(info.get(key, ""))
                    if val:
                        items.append(MetadataItem(self._generate_id(), file_path, fname, key, val,
                                                  "high" if fname in ("author", "creator") else "medium",
                                                  f"Review {fname} metadata before distribution"))
            except Exception:
                pass
        elif ext in (".xlsx", ".pptx") and HAS_OPENPYXL:
            try:
                wb = load_workbook(file_path, read_only=True)
                props = wb.properties
                for fname, key, val, risk in [
                    ("author", "author", getattr(props, "creator", ""), "high"),
                    ("created_date", "created_date", str(getattr(props, "created", "")), "medium"),
                    ("modified_date", "modified_date", str(getattr(props, "modified", "")), "medium"),
                    ("last_modified_by", "last_modified_by", getattr(props, "lastModifiedBy", ""), "medium"),
                ]:
                    if val:
                        items.append(MetadataItem(self._generate_id(), file_path, fname, key, str(val), risk,
                                                  f"Review {fname} metadata before sharing"))
                wb.close()
            except Exception:
                pass
        return items

    def _detect_hidden_content(self, file_path: str) -> Tuple[bool, List[str]]:
        issues: List[str] = []
        ext = os.path.splitext(file_path)[1].lower()
        try:
            if ext == ".xlsx" and HAS_OPENPYXL:
                wb = load_workbook(file_path, read_only=True)
                for ws in wb.worksheets:
                    if ws.sheet_state != "visible":
                        issues.append(f"Hidden sheet: {ws.title}")
                    for col, dim in ws.column_dimensions.items():
                        if dim.hidden:
                            issues.append(f"Hidden column: {col}")
                    for row, dim in ws.row_dimensions.items():
                        if dim.hidden:
                            issues.append(f"Hidden row: {row}")
                wb.close()
            elif ext == ".docx" and HAS_DOCX:
                for para in DocxDocument(file_path).paragraphs:
                    for run in para.runs:
                        color = str(run.font.color.rgb) if run.font.color and run.font.color.rgb else ""
                        if color in ("FFFFFFFF", "FFFFFF00") and run.text.strip():
                            issues.append("Hidden text with white/yellow font")
                            break
        except Exception:
            pass
        return len(issues) > 0, issues

    def _scan_word_doc(self, fp: str) -> Tuple[str, List[MetadataItem]]:
        if not HAS_DOCX:
            return "", []
        try:
            doc = DocxDocument(fp)
            parts = [p.text for p in doc.paragraphs if p.text.strip()]
            for t in doc.tables:
                for r in t.rows:
                    parts.extend(c.text for c in r.cells if c.text.strip())
            return "\n".join(parts), []
        except Exception:
            return "", []

    def _scan_pdf(self, fp: str) -> Tuple[str, List[MetadataItem]]:
        if not HAS_PYPDF:
            return "", []
        try:
            parts = [p.extract_text() or "" for p in PdfReader(fp).pages]
            return "\n".join(parts), []
        except Exception:
            return "", []

    def _scan_excel(self, fp: str) -> Tuple[str, List[MetadataItem]]:
        if not HAS_OPENPYXL:
            return "", []
        try:
            wb = load_workbook(fp, read_only=True, data_only=True)
            parts = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    parts.extend(str(v) for v in row if v is not None)
            wb.close()
            return "\n".join(parts), []
        except Exception:
            return "", []

    def _scan_powerpoint(self, fp: str) -> Tuple[str, List[MetadataItem]]:
        try:
            if not zipfile.is_zipfile(fp):
                return "", []
            texts: List[str] = []
            with zipfile.ZipFile(fp, "r") as zf:
                for name in zf.namelist():
                    if name.startswith("ppt/slides/slide") and name.endswith(".xml"):
                        data = zf.read(name).decode("utf-8", errors="ignore")
                        texts.extend(t for t in re.findall(r"<a:t[^>]*>([^<]+)</a:t>", data) if t.strip())
            return "\n".join(texts), []
        except Exception:
            return "", []

    def _scan_text_file(self, fp: str) -> Tuple[str, List[MetadataItem]]:
        for enc in ("utf-8", "latin-1", "ascii"):
            try:
                with open(fp, "r", encoding=enc, errors="strict") as f:
                    return f.read(), []
            except (UnicodeDecodeError, UnicodeError):
                continue
        try:
            with open(fp, "r", encoding="utf-8", errors="ignore") as f:
                return f.read(), []
        except Exception:
            return "", []

    def _scan_csv(self, fp: str) -> Tuple[str, List[MetadataItem]]:
        for enc in ("utf-8", "latin-1", "ascii"):
            try:
                with open(fp, "r", encoding=enc, errors="strict", newline="") as f:
                    return "\n".join(",".join(r) for r in csv.reader(f)), []
            except (UnicodeDecodeError, UnicodeError):
                continue
        return "", []

    def scan_file(self, file_path: str) -> SecurityReport:
        with self._lock:
            ext = os.path.splitext(file_path)[1].lower()
            scanner_map = {
                ".docx": self._scan_word_doc, ".pdf": self._scan_pdf,
                ".xlsx": self._scan_excel, ".xls": self._scan_excel,
                ".pptx": self._scan_powerpoint, ".ppt": self._scan_powerpoint,
                ".csv": self._scan_csv, ".txt": self._scan_text_file,
                ".md": self._scan_text_file, ".json": self._scan_text_file,
                ".xml": self._scan_text_file, ".yaml": self._scan_text_file,
                ".yml": self._scan_text_file, ".ini": self._scan_text_file,
                ".cfg": self._scan_text_file, ".conf": self._scan_text_file,
                ".py": self._scan_text_file, ".js": self._scan_text_file,
                ".ts": self._scan_text_file, ".java": self._scan_text_file,
                ".c": self._scan_text_file, ".cpp": self._scan_text_file,
                ".h": self._scan_text_file, ".sh": self._scan_text_file,
                ".bat": self._scan_text_file, ".ps1": self._scan_text_file,
                ".log": self._scan_text_file,
            }
            content, extra_meta = scanner_map.get(ext, self._scan_text_file)(file_path)
            findings = self._scan_content(file_path, content) if content else []
            metadata = self._scan_metadata(file_path) + extra_meta
            hidden_found, _ = self._detect_hidden_content(file_path)
            risk_score = self._compute_risk_score(findings, metadata)
            if hidden_found:
                risk_score = min(100.0, risk_score + 10.0)
            overall_risk = self._assess_risk(findings, metadata)
            recommendations = self._generate_recommendations(findings, metadata)
            if hidden_found:
                recommendations.append("Hidden content detected - review and remove before sharing")
            report = SecurityReport(
                report_id=self._generate_id(), file_path=file_path, scan_timestamp=time.time(),
                overall_risk=overall_risk, findings=findings, metadata_items=metadata,
                hidden_content_found=hidden_found, total_findings=len(findings),
                risk_score=risk_score, recommendations=recommendations,
            )
            self._scan_history.append({
                "report_id": report.report_id, "file_path": file_path,
                "timestamp": report.scan_timestamp, "risk_score": risk_score,
                "total_findings": len(findings),
            })
            if len(self._scan_history) > self._max_history:
                self._scan_history = self._scan_history[-self._max_history:]
            return report

    def _assess_risk(self, findings: List[SensitiveFinding], metadata: List[MetadataItem]) -> SensitivityLevel:
        score = self._compute_risk_score(findings, metadata)
        if score >= 80:
            return SensitivityLevel.TOP_SECRET
        elif score >= 60:
            return SensitivityLevel.SECRET
        elif score >= 40:
            return SensitivityLevel.CONFIDENTIAL
        elif score >= 20:
            return SensitivityLevel.INTERNAL
        return SensitivityLevel.PUBLIC

    def _compute_risk_score(self, findings: List[SensitiveFinding], metadata: List[MetadataItem]) -> float:
        if not findings and not metadata:
            return 0.0
        weights = {
            SensitiveCategory.CREDENTIAL: 30.0, SensitiveCategory.PII: 25.0,
            SensitiveCategory.FINANCIAL: 25.0, SensitiveCategory.MEDICAL: 25.0,
            SensitiveCategory.CLASSIFICATION: 20.0, SensitiveCategory.PROPRIETARY: 15.0,
            SensitiveCategory.CONTACT: 10.0, SensitiveCategory.LOCATION: 8.0,
        }
        score = sum(weights.get(f.category, 5.0) * f.confidence for f in findings)
        score += min(len(set(f.category for f in findings)) * 3.0, 15.0)
        score += sum(10.0 if m.risk == "high" else 5.0 if m.risk == "medium" else 1.0 for m in metadata)
        return min(100.0, round(score, 2))

    def _generate_recommendations(self, findings: List[SensitiveFinding], metadata: List[MetadataItem]) -> List[str]:
        recs: List[str] = []
        if not findings and not metadata:
            return ["No sensitive content detected"]
        cats = set(f.category for f in findings)
        rec_map = {
            SensitiveCategory.CREDENTIAL: ["CRITICAL: Remove all credentials and rotate them immediately",
                                            "Implement secrets management for sensitive configuration"],
            SensitiveCategory.PII: ["Review PII exposure and ensure GDPR/CCPA compliance",
                                    "Consider data anonymization or masking for PII fields"],
            SensitiveCategory.FINANCIAL: ["Remove financial data and ensure PCI DSS compliance"],
            SensitiveCategory.MEDICAL: ["Remove medical information and ensure HIPAA compliance"],
            SensitiveCategory.CLASSIFICATION: ["Review document classification markings and handle accordingly"],
            SensitiveCategory.PROPRIETARY: ["Review proprietary information disclosure risks"],
            SensitiveCategory.CONTACT: ["Review contact information for privacy exposure"],
            SensitiveCategory.LOCATION: ["Review location data for privacy concerns"],
        }
        for cat in cats:
            recs.extend(rec_map.get(cat, []))
        high_meta = [m for m in metadata if m.risk == "high"]
        if high_meta:
            recs.append(f"Remove high-risk metadata ({len(high_meta)} items found)")
        if any(m.risk == "medium" for m in metadata):
            recs.append("Strip metadata before sharing documents externally")
        return recs

    def scrub_file(self, file_path: str, output_path: Optional[str] = None,
                   categories: Optional[List[SensitiveCategory]] = None) -> Dict[str, Any]:
        with self._lock:
            if not os.path.isfile(file_path):
                return {"success": False, "error": "File not found"}
            if not output_path:
                base, ext = os.path.splitext(file_path)
                output_path = f"{base}_scrubbed{ext}"
            try:
                with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                    content = f.read()
            except Exception as e:
                return {"success": False, "error": str(e)}
            all_patterns = {**self._builtin_patterns, **self._custom_patterns}
            target = {n: c for n, c in all_patterns.items() if categories is None or c["category"] in categories}
            count = 0
            for cfg in target.values():
                try:
                    compiled = re.compile(cfg["pattern"], re.IGNORECASE)
                    content, n = compiled.subn(lambda _: "[REDACTED]", content)
                    count += n
                except re.error:
                    continue
            try:
                with open(output_path, "w", encoding="utf-8") as f:
                    f.write(content)
            except Exception as e:
                return {"success": False, "error": str(e)}
            return {"success": True, "output_path": output_path, "replacements_made": count}

    def add_custom_pattern(self, name: str, pattern: str, category: SensitiveCategory,
                           recommendation: str) -> None:
        with self._lock:
            self._custom_patterns[name] = {
                "pattern": pattern, "category": category, "confidence": 0.7,
                "recommendation": recommendation,
            }

    def batch_scan(self, file_paths: List[str]) -> List[SecurityReport]:
        return [self.scan_file(fp) for fp in file_paths if os.path.isfile(fp)]

    def get_scan_history(self, limit: int = 50) -> List[Dict[str, Any]]:
        with self._lock:
            return list(self._scan_history[-limit:])

    def get_scanner_summary(self) -> Dict[str, Any]:
        with self._lock:
            cats = list({c["category"].value for c in self._builtin_patterns.values()})
            return {
                "version": self.VERSION, "builtin_patterns": len(self._builtin_patterns),
                "custom_patterns": len(self._custom_patterns),
                "total_patterns": len(self._builtin_patterns) + len(self._custom_patterns),
                "categories_covered": cats, "total_scans_performed": len(self._scan_history),
                "has_docx_support": HAS_DOCX, "has_pdf_support": HAS_PYPDF,
                "has_xlsx_support": HAS_OPENPYXL, "max_history": self._max_history,
            }
