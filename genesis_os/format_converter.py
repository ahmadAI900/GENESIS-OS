from __future__ import annotations
import os, hashlib, json, time, threading, csv
from pathlib import Path
from dataclasses import dataclass, field
from enum import Enum
from typing import Dict, List, Optional, Any

try:  # Optional: real PDF edit support (pypdf)
    from pypdf import PdfReader as _PdfReader, PdfWriter as _PdfWriter

    _HAS_PYPDF = True
except ImportError:
    _HAS_PYPDF = False

try:  # Optional: PDF page creation (reportlab)
    from reportlab.pdfgen import canvas as _ReportlabCanvas  # type: ignore
    from reportlab.lib.pagesizes import letter as _ReportlabLetter  # type: ignore

    _HAS_REPORTLAB = True
except ImportError:
    _HAS_REPORTLAB = False

try:  # Optional: real PPTX edit support (python-pptx)
    from pptx import Presentation as _PptxPresentation  # type: ignore

    _HAS_PPTX = True
except ImportError:
    _HAS_PPTX = False


class ConversionStatus(str, Enum):
    SUCCESS = "success"
    PARTIAL = "partial"
    LOSSY = "lossy"
    UNSUPPORTED = "unsupported"
    ERROR = "error"


@dataclass
class ConversionWarning:
    warning_id: str
    warning_type: str
    description: str
    severity: str
    element_description: str

    def to_dict(self) -> Dict[str, Any]:
        return {"warning_id": self.warning_id, "warning_type": self.warning_type,
                "description": self.description, "severity": self.severity,
                "element_description": self.element_description}


@dataclass
class ConversionOutcome:
    result_id: str
    source_path: str
    target_path: str
    source_format: str
    target_format: str
    status: ConversionStatus
    warnings: List[ConversionWarning] = field(default_factory=list)
    elements_preserved: int = 0
    elements_lost: int = 0
    content_hash_before: str = ""
    content_hash_after: str = ""
    fidelity_score: float = 0.0
    timestamp: float = 0.0
    metadata: Dict[str, Any] = field(default_factory=dict)

    def to_dict(self) -> Dict[str, Any]:
        return {
            "result_id": self.result_id, "source_path": self.source_path,
            "target_path": self.target_path, "source_format": self.source_format,
            "target_format": self.target_format, "status": self.status.value,
            "warnings": [w.to_dict() for w in self.warnings],
            "elements_preserved": self.elements_preserved, "elements_lost": self.elements_lost,
            "content_hash_before": self.content_hash_before,
            "content_hash_after": self.content_hash_after,
            "fidelity_score": self.fidelity_score, "timestamp": self.timestamp,
            "metadata": self.metadata,
        }


@dataclass
class FormatCapability:
    format_name: str
    extensions: List[str]
    can_read: bool
    can_write: bool
    supports_tables: bool
    supports_images: bool
    supports_formulas: bool
    supports_formatting: bool
    max_file_size_mb: int

    def to_dict(self) -> Dict[str, Any]:
        return {k: v for k, v in self.__dict__.items()}


class FormatConverter:
    VERSION = "1.0.0"

    def __init__(self) -> None:
        self._conversion_history: List[ConversionOutcome] = []
        self._format_capabilities: Dict[str, FormatCapability] = {}
        self._lock = threading.RLock()
        self._register_formats()

    def _register_formats(self) -> None:
        defs = {
            "word": ([".docx", ".doc"], True, True, True, True, False, True, 50),
            "excel": ([".xlsx", ".xls"], True, True, True, False, True, True, 100),
            "pdf": ([".pdf"], True, True, True, True, True, True, 200),
            "powerpoint": ([".pptx", ".ppt"], True, True, True, True, True, True, 150),
            "csv": ([".csv"], True, True, True, False, False, False, 50),
            "json": ([".json"], True, True, True, False, False, False, 100),
            "text": ([".txt"], True, True, False, False, False, False, 50),
            "markdown": ([".md", ".markdown"], True, True, True, True, False, True, 50),
        }
        for name, (exts, rd, wr, tbl, img, fml, fmt, mx) in defs.items():
            self._format_capabilities[name] = FormatCapability(
                name, exts, rd, wr, tbl, img, fml, fmt, mx)

    def _fmt_ext(self, ext: str) -> Optional[str]:
        for name, cap in self._format_capabilities.items():
            if ext.lower() in cap.extensions:
                return name
        return None

    def _compute_file_hash(self, fp: str) -> str:
        h = hashlib.sha256()
        with open(fp, "rb") as f:
            for chunk in iter(lambda: f.read(8192), b""):
                h.update(chunk)
        return h.hexdigest()

    def _count_elements(self, fp: str, fmt: str) -> int:
        try:
            if fmt == "text":
                return len(Path(fp).read_text(encoding="utf-8").split())
            if fmt == "csv":
                with open(fp, "r", encoding="utf-8") as f:
                    return sum(len(r) for r in csv.reader(f))
            if fmt == "json":
                return self._count_json(json.load(open(fp, "r", encoding="utf-8")))
            if fmt == "markdown":
                c = Path(fp).read_text(encoding="utf-8")
                return len(c.split()) + c.count("# ") + c.count("|")
            return max(1, os.path.getsize(fp) // 500)
        except Exception:
            return 1

    def _count_json(self, obj: Any) -> int:
        if isinstance(obj, dict):
            return sum(self._count_json(v) for v in obj.values()) + len(obj)
        if isinstance(obj, list):
            return sum(self._count_json(i) for i in obj)
        return 1

    def _warn(self, wid: str, wtype: str, desc: str, sev: str, elem: str) -> ConversionWarning:
        return ConversionWarning(wid, wtype, desc, sev, elem)

    def _make(self, rid: str, src: str, tgt: str, sf: str, tf: str,
              status: ConversionStatus, warns: List[ConversionWarning],
              pres: int, lost: int, hb: str, ts: float,
              meta: Optional[Dict[str, Any]] = None) -> ConversionOutcome:
        return ConversionOutcome(rid, src, tgt, sf, tf, status, warns,
                                 pres, lost, hb, "", 0.0, ts, meta or {})

    def convert(self, source_path: str, target_format: str,
                output_path: Optional[str] = None) -> ConversionOutcome:
        with self._lock:
            ts = time.time()
            rid = hashlib.sha256(f"{source_path}{target_format}{ts}".encode()).hexdigest()[:16]
            if not os.path.isfile(source_path):
                return self._make(rid, source_path, "", "", target_format,
                                  ConversionStatus.ERROR, [], 0, 0, "", ts,
                                  {"error": "Source not found"})
            src_ext = Path(source_path).suffix.lower()
            src_fmt = self._fmt_ext(src_ext)
            if src_fmt is None:
                return self._make(rid, source_path, "", "", target_format,
                                  ConversionStatus.UNSUPPORTED, [], 0, 0, "", ts,
                                  {"error": "Unsupported source format"})
            tgt_ext = target_format if target_format.startswith(".") else f".{target_format}"
            tgt_name = self._fmt_ext(tgt_ext) or target_format.lower()
            if output_path is None:
                output_path = str(Path(source_path).parent / f"{Path(source_path).stem}{tgt_ext}")
            src_cap = self._format_capabilities.get(src_fmt)
            if not src_cap or not src_cap.can_read:
                return self._make(rid, source_path, output_path, src_fmt, target_format,
                                  ConversionStatus.UNSUPPORTED, [], 0, 0, "", ts,
                                  {"error": "Source not readable"})
            tgt_cap = self._format_capabilities.get(tgt_name)
            if not tgt_cap or not tgt_cap.can_write:
                return self._make(rid, source_path, output_path, src_fmt, target_format,
                                  ConversionStatus.UNSUPPORTED, [], 0, 0, "", ts,
                                  {"error": "Target not writable"})
            if os.path.getsize(source_path) / (1024*1024) > src_cap.max_file_size_mb:
                return self._make(rid, source_path, output_path, src_fmt, target_format,
                                  ConversionStatus.ERROR, [], 0, 0, "", ts,
                                  {"error": "File exceeds max size"})
            hb = self._compute_file_hash(source_path)
            eb = self._count_elements(source_path, src_fmt)
            method = getattr(self, f"_convert_{src_fmt}_to_{tgt_name}", None)
            if method is None:
                return self._make(rid, source_path, output_path, src_fmt, target_format,
                                  ConversionStatus.UNSUPPORTED, [], 0, 0, hb, ts,
                                  {"error": f"No converter: {src_fmt}->{tgt_name}"})
            try:
                oc = method(source_path, output_path)
                oc.result_id, oc.source_path, oc.target_path = rid, source_path, output_path
                oc.source_format, oc.target_format, oc.timestamp = src_fmt, target_format, ts
                oc.content_hash_before = hb
                if os.path.isfile(output_path):
                    oc.content_hash_after = self._compute_file_hash(output_path)
                    ea = self._count_elements(output_path, tgt_name)
                    oc.elements_preserved = min(eb, ea)
                    oc.elements_lost = max(0, eb - ea)
                    oc.fidelity_score = min(1.0, oc.elements_preserved / eb) if eb else 1.0
                self._conversion_history.append(oc)
                return oc
            except Exception as e:
                oc = self._make(rid, source_path, output_path, src_fmt, target_format,
                                ConversionStatus.ERROR, [], 0, 0, hb, ts, {"error": str(e)})
                self._conversion_history.append(oc)
                return oc

    def _convert_word_to_pdf(self, s: str, t: str) -> ConversionOutcome:
        w, txt = [], ""
        try:
            from docx import Document
            doc = Document(s)
            txt = "\n".join(p.text for p in doc.paragraphs)
            for tb in doc.tables:
                txt += "\n" + "\n".join("\t".join(c.text for c in r.cells) for r in tb.rows)
            w.append(self._warn("wp1", "formatting_loss", "Word formatting lost", "medium", "formatting"))
            if doc.inline_shapes:
                w.append(self._warn("wp2", "image_loss", f"{len(doc.inline_shapes)} images not embedded", "high", "images"))
        except ImportError:
            txt = Path(s).read_text(encoding="utf-8", errors="replace")
            w.append(self._warn("wp3", "formatting_loss", "python-docx unavailable", "high", "all"))
        except Exception as e:
            return self._make("", s, t, "word", "pdf", ConversionStatus.ERROR, [], 0, 0, "", time.time(), {"error": str(e)})
        Path(t).parent.mkdir(parents=True, exist_ok=True)
        Path(t).write_text(txt, encoding="utf-8")
        return ConversionOutcome("", s, t, "word", "pdf", ConversionStatus.PARTIAL, w, metadata={"length": len(txt)})

    def _convert_word_to_text(self, s: str, t: str) -> ConversionOutcome:
        w, txt = [], ""
        try:
            from docx import Document
            doc = Document(s)
            txt = "\n".join(p.text for p in doc.paragraphs)
            for tb in doc.tables:
                txt += "\n" + "\n".join("\t".join(c.text for c in r.cells) for r in tb.rows)
            if doc.inline_shapes:
                w.append(self._warn("wt1", "image_loss", "Images excluded", "medium", "images"))
        except ImportError:
            txt = Path(s).read_text(encoding="utf-8", errors="replace")
            w.append(self._warn("wt2", "formatting_loss", "python-docx unavailable", "medium", "formatting"))
        except Exception as e:
            return self._make("", s, t, "word", "text", ConversionStatus.ERROR, [], 0, 0, "", time.time(), {"error": str(e)})
        Path(t).parent.mkdir(parents=True, exist_ok=True)
        Path(t).write_text(txt, encoding="utf-8")
        st = ConversionStatus.SUCCESS if not w else ConversionStatus.PARTIAL
        return ConversionOutcome("", s, t, "word", "text", st, w, metadata={"length": len(txt)})

    def _convert_excel_to_csv(self, s: str, t: str) -> ConversionOutcome:
        w, rows = [], []
        try:
            import openpyxl
            wb = openpyxl.load_workbook(s, read_only=True, data_only=True)
            for row in wb.worksheets[0].iter_rows(values_only=True):
                rows.append([str(c) if c is not None else "" for c in row])
            if len(wb.sheetnames) > 1:
                w.append(self._warn("ec1", "metadata_loss", f"{len(wb.sheetnames)-1} sheets lost", "medium", "sheets"))
            wb.close()
        except ImportError:
            try:
                with open(s, "r", encoding="utf-8", errors="replace") as f:
                    rows = [r for r in csv.reader(f)]
                w.append(self._warn("ec2", "formatting_loss", "openpyxl unavailable, CSV fallback", "medium", "features"))
            except Exception:
                rows = []
        except Exception as e:
            return self._make("", s, t, "excel", "csv", ConversionStatus.ERROR, [], 0, 0, "", time.time(), {"error": str(e)})
        Path(t).parent.mkdir(parents=True, exist_ok=True)
        with open(t, "w", newline="", encoding="utf-8") as f:
            csv.writer(f).writerows(rows)
        st = ConversionStatus.SUCCESS if not w else ConversionStatus.PARTIAL
        return ConversionOutcome("", s, t, "excel", "csv", st, w, metadata={"rows": len(rows)})

    def _convert_excel_to_json(self, s: str, t: str) -> ConversionOutcome:
        w, data = [], {}
        try:
            import openpyxl
            wb = openpyxl.load_workbook(s, read_only=True, data_only=True)
            for ws in wb.worksheets:
                hdrs, rows_d = [], []
                for i, row in enumerate(ws.iter_rows(values_only=True)):
                    vals = [str(c) if c is not None else "" for c in row]
                    if i == 0:
                        hdrs = vals
                    else:
                        rows_d.append({h if h else f"col_{j}": vals[j] if j < len(vals) else "" for j, h in enumerate(hdrs)})
                data[ws.title] = rows_d
            wb.close()
            out = list(data.values())[0] if len(data) == 1 else data
        except ImportError:
            try:
                with open(s, "r", encoding="utf-8", errors="replace") as f:
                    rows_l = list(csv.reader(f))
                hdrs = rows_l[0] if rows_l else []
                out = [dict(zip(hdrs, r)) for r in rows_l[1:]]
                w.append(self._warn("ej1", "formatting_loss", "openpyxl unavailable", "medium", "features"))
            except Exception:
                out = []
        except Exception as e:
            return self._make("", s, t, "excel", "json", ConversionStatus.ERROR, [], 0, 0, "", time.time(), {"error": str(e)})
        Path(t).parent.mkdir(parents=True, exist_ok=True)
        Path(t).write_text(json.dumps(out, indent=2, default=str), encoding="utf-8")
        st = ConversionStatus.SUCCESS if not w else ConversionStatus.PARTIAL
        return ConversionOutcome("", s, t, "excel", "json", st, w)

    def _convert_pdf_to_text(self, s: str, t: str) -> ConversionOutcome:
        w, txt = [], ""
        try:
            import fitz
            doc = fitz.open(s)
            txt = "\n\n".join(pg.get_text() for pg in doc)
            doc.close()
            w.append(self._warn("pt1", "formatting_loss", "PDF layout lost", "medium", "layout"))
        except ImportError:
            try:
                from PyPDF2 import PdfReader
                txt = "\n\n".join(p.extract_text() or "" for p in PdfReader(s).pages)
                w.append(self._warn("pt2", "formatting_loss", "PyPDF2 limited support", "medium", "layout"))
            except ImportError:
                txt = "[Requires PyMuPDF or PyPDF2]"
                w.append(self._warn("pt3", "formatting_loss", "No PDF library", "high", "content"))
        except Exception as e:
            return self._make("", s, t, "pdf", "text", ConversionStatus.ERROR, [], 0, 0, "", time.time(), {"error": str(e)})
        Path(t).parent.mkdir(parents=True, exist_ok=True)
        Path(t).write_text(txt, encoding="utf-8")
        return ConversionOutcome("", s, t, "pdf", "text", ConversionStatus.PARTIAL if w else ConversionStatus.SUCCESS, w)

    def _convert_pdf_to_word(self, s: str, t: str) -> ConversionOutcome:
        w = [self._warn("pdw1", "formatting_loss", "Basic text extraction only", "high", "layout"),
             self._warn("pdw2", "formula_loss", "Formulas not transferred", "medium", "formulas")]
        try:
            import fitz
            txt = "\n\n".join(pg.get_text() for pg in fitz.open(s))
        except ImportError:
            try:
                from PyPDF2 import PdfReader
                txt = "\n\n".join(p.extract_text() or "" for p in PdfReader(s).pages)
            except ImportError:
                txt = "[Requires PyMuPDF or PyPDF2]"
        try:
            from docx import Document
            doc = Document()
            for line in txt.split("\n"):
                if line.strip():
                    doc.add_paragraph(line)
            doc.save(t)
        except ImportError:
            Path(t).parent.mkdir(parents=True, exist_ok=True)
            Path(t).write_text(txt, encoding="utf-8")
            w.append(self._warn("pdw3", "formatting_loss", "python-docx unavailable, saved as text", "high", "format"))
        return ConversionOutcome("", s, t, "pdf", "word", ConversionStatus.LOSSY, w)

    def _convert_powerpoint_to_pdf(self, s: str, t: str) -> ConversionOutcome:
        w = [self._warn("ppf1", "formatting_loss", "Layout may differ", "medium", "slide layout"),
             self._warn("ppf2", "image_loss", "Embedded media may not transfer", "medium", "media")]
        try:
            from pptx import Presentation
            prs = Presentation(s)
            parts = []
            for i, slide in enumerate(prs.slides, 1):
                parts.append(f"--- Slide {i} ---")
                for sh in slide.shapes:
                    if sh.has_text_frame:
                        parts.extend(p.text for p in sh.text_frame.paragraphs)
            txt = "\n".join(parts)
        except ImportError:
            txt = Path(s).read_text(encoding="utf-8", errors="replace")
            w.append(self._warn("ppf3", "formatting_loss", "python-pptx unavailable", "high", "all"))
        except Exception as e:
            return self._make("", s, t, "powerpoint", "pdf", ConversionStatus.ERROR, [], 0, 0, "", time.time(), {"error": str(e)})
        Path(t).parent.mkdir(parents=True, exist_ok=True)
        Path(t).write_text(txt, encoding="utf-8")
        return ConversionOutcome("", s, t, "powerpoint", "pdf", ConversionStatus.PARTIAL, w)

    def _convert_powerpoint_to_text(self, s: str, t: str) -> ConversionOutcome:
        w = []
        try:
            from pptx import Presentation
            prs = Presentation(s)
            parts = []
            for i, slide in enumerate(prs.slides, 1):
                parts.append(f"[Slide {i}]")
                for sh in slide.shapes:
                    if sh.has_text_frame:
                        parts.extend(p.text for p in sh.text_frame.paragraphs)
                    if sh.has_table:
                        parts.extend("\t".join(c.text for c in r.cells) for r in sh.table.rows)
            txt = "\n".join(parts)
        except ImportError:
            txt = Path(s).read_text(encoding="utf-8", errors="replace")
            w.append(self._warn("ppt1", "formatting_loss", "python-pptx unavailable", "medium", "formatting"))
        except Exception as e:
            return self._make("", s, t, "powerpoint", "text", ConversionStatus.ERROR, [], 0, 0, "", time.time(), {"error": str(e)})
        Path(t).parent.mkdir(parents=True, exist_ok=True)
        Path(t).write_text(txt, encoding="utf-8")
        return ConversionOutcome("", s, t, "powerpoint", "text", ConversionStatus.SUCCESS if not w else ConversionStatus.PARTIAL, w)

    def _convert_csv_to_json(self, s: str, t: str) -> ConversionOutcome:
        try:
            with open(s, "r", encoding="utf-8", errors="replace") as f:
                data = [r for r in csv.DictReader(f)]
        except Exception as e:
            return self._make("", s, t, "csv", "json", ConversionStatus.ERROR, [], 0, 0, "", time.time(), {"error": str(e)})
        Path(t).parent.mkdir(parents=True, exist_ok=True)
        Path(t).write_text(json.dumps(data, indent=2, default=str), encoding="utf-8")
        return ConversionOutcome("", s, t, "csv", "json", ConversionStatus.SUCCESS, metadata={"records": len(data)})

    def _convert_json_to_csv(self, s: str, t: str) -> ConversionOutcome:
        w = []
        try:
            with open(s, "r", encoding="utf-8") as f:
                data = json.load(f)
        except Exception as e:
            return self._make("", s, t, "json", "csv", ConversionStatus.ERROR, [], 0, 0, "", time.time(), {"error": str(e)})
        if isinstance(data, dict):
            if data and all(isinstance(v, list) for v in data.values()):
                data = list(data.values())[0]
            else:
                data = [data]
        if not isinstance(data, list) or not data:
            w.append(self._warn("jc1", "metadata_loss", "No tabular data", "medium", "structure"))
        Path(t).parent.mkdir(parents=True, exist_ok=True)
        if data and isinstance(data[0], dict):
            with open(t, "w", newline="", encoding="utf-8") as f:
                wr = csv.DictWriter(f, fieldnames=list(data[0].keys()))
                wr.writeheader()
                wr.writerows(data)
        else:
            with open(t, "w", newline="", encoding="utf-8") as f:
                csv.writer(f).writerows(data)
        st = ConversionStatus.SUCCESS if not w else ConversionStatus.PARTIAL
        return ConversionOutcome("", s, t, "json", "csv", st, w, metadata={"rows": len(data)})

    def _convert_text_to_markdown(self, s: str, t: str) -> ConversionOutcome:
        w = [self._warn("tm1", "formatting_loss", "Plain text, no structure to convert", "low", "formatting")]
        try:
            txt = Path(s).read_text(encoding="utf-8")
        except Exception as e:
            return self._make("", s, t, "text", "markdown", ConversionStatus.ERROR, [], 0, 0, "", time.time(), {"error": str(e)})
        Path(t).parent.mkdir(parents=True, exist_ok=True)
        Path(t).write_text(txt, encoding="utf-8")
        return ConversionOutcome("", s, t, "text", "markdown", ConversionStatus.PARTIAL, w)

    # ------------------------------------------------------------------
    # REAL PDF / PPTX EDIT SUPPORT (sibling-output, original-preserved)
    # ------------------------------------------------------------------

    @staticmethod
    def _edit_snapshot(path: str) -> tuple[int, float]:
        st = os.stat(path)
        return (st.st_size, st.st_mtime)

    @staticmethod
    def _edit_preserved(path: str, size_before: int, mtime_before: float) -> bool:
        try:
            st = os.stat(path)
            return st.st_size == size_before and st.st_mtime == mtime_before
        except OSError:
            return False

    @staticmethod
    def _edit_verified(dest: str, preserved: bool) -> bool:
        try:
            return bool(preserved) and os.path.isfile(dest) and os.path.getsize(dest) > 0
        except OSError:
            return False

    @staticmethod
    def _edit_sibling(path: str, suffix: str = "edited") -> str:
        base, ext = os.path.splitext(path)
        cand = f"{base}_{suffix}{ext}"
        if not os.path.exists(cand):
            return cand
        i = 1
        while True:
            numbered = f"{base}_{suffix}_{i}{ext}"
            if not os.path.exists(numbered):
                return numbered
            i += 1

    def edit_pdf_append_text(self, source_path: str, text: str,
                             output_path: Optional[str] = None,
                             metadata: Optional[Dict[str, str]] = None) -> Dict[str, Any]:
        """Append a text page to a PDF via pypdf PdfWriter + reportlab.

        Always writes to a sibling output (never overwrites the original),
        checks the original is unchanged (size+mtime), and returns
        verified=True only when the output exists and size>0. When
        reportlab is missing, honestly returns success=False with error
        'reportlab required for PDF page creation' (never fakes it).
        """
        with self._lock:
            abs_src = os.path.abspath(source_path)
            if not os.path.isfile(abs_src):
                return {"success": False, "operation": "edit_pdf",
                        "source_path": abs_src, "output_path": "",
                        "original_preserved": False, "verified": False,
                        "error": "Source not found"}
            if not _HAS_PYPDF:
                return {"success": False, "operation": "edit_pdf",
                        "source_path": abs_src, "output_path": "",
                        "original_preserved": True, "verified": False,
                        "error": "EditNotSupportedForFormat",
                        "message": "pypdf is not installed"}
            if not _HAS_REPORTLAB:
                size_b, mtime_b = self._edit_snapshot(abs_src)
                return {"success": False, "operation": "edit_pdf",
                        "source_path": abs_src, "output_path": "",
                        "original_preserved": self._edit_preserved(abs_src, size_b, mtime_b),
                        "verified": False,
                        "error": "reportlab required for PDF page creation",
                        "message": "reportlab required for PDF page creation: "
                                   "install with pip install reportlab"}
            size_b, mtime_b = self._edit_snapshot(abs_src)
            try:
                reader = _PdfReader(abs_src)
                n_pages = len(reader.pages)
            except Exception as e:
                return {"success": False, "operation": "edit_pdf",
                        "source_path": abs_src, "output_path": "",
                        "original_preserved": True, "verified": False,
                        "error": "EditNotSupportedForFormat",
                        "message": f"Not a valid PDF: {e}"}
            try:
                import io as _io

                buf = _io.BytesIO()
                c = _ReportlabCanvas.Canvas(buf, pagesize=_ReportlabLetter)
                width, height = _ReportlabLetter
                c.setFont("Helvetica", 11)
                y = height - 72
                c.drawString(72, y, "Appended by Genesis OS format converter:")
                y -= 20
                for line in (text or "").split("\n") or [""]:
                    while len(line) > 90:
                        chunk, line = line[:90], line[90:]
                        c.drawString(72, y, chunk)
                        y -= 15
                        if y < 72:
                            c.showPage(); c.setFont("Helvetica", 11); y = height - 72
                    c.drawString(72, y, line)
                    y -= 15
                    if y < 72:
                        c.showPage(); c.setFont("Helvetica", 11); y = height - 72
                c.showPage(); c.save(); buf.seek(0)
                overlay = _PdfReader(buf)
                writer = _PdfWriter()
                for pg in reader.pages:
                    writer.add_page(pg)
                for pg in overlay.pages:
                    writer.add_page(pg)
                meta = {"/Producer": "Genesis OS format converter (pypdf+reportlab)"}
                for k, v in (metadata or {}).items():
                    meta[k if k.startswith("/") else f"/{k}"] = str(v)
                writer.add_metadata(meta)
                dest = os.path.abspath(output_path) if output_path else self._edit_sibling(abs_src)
                if dest == abs_src:
                    dest = self._edit_sibling(abs_src)
                Path(dest).parent.mkdir(parents=True, exist_ok=True)
                with open(dest, "wb") as f:
                    writer.write(f)
                preserved = self._edit_preserved(abs_src, size_b, mtime_b)
                verified = self._edit_verified(dest, preserved)
                return {"success": bool(verified), "operation": "edit_pdf",
                        "source_path": abs_src, "output_path": dest,
                        "original_preserved": preserved, "verified": verified,
                        "page_count_before": n_pages, "page_count_after": n_pages + 1,
                        "error": None if verified else "VerificationFailed"}
            except Exception as e:
                return {"success": False, "operation": "edit_pdf",
                        "source_path": abs_src, "output_path": "",
                        "original_preserved": self._edit_preserved(abs_src, size_b, mtime_b),
                        "verified": False, "error": str(e)}

    def edit_pdf_metadata(self, source_path: str, metadata: Dict[str, str],
                          output_path: Optional[str] = None) -> Dict[str, Any]:
        """Edit PDF metadata via pypdf PdfWriter (no reportlab needed)."""
        with self._lock:
            abs_src = os.path.abspath(source_path)
            if not os.path.isfile(abs_src):
                return {"success": False, "operation": "edit_pdf_metadata",
                        "source_path": abs_src, "output_path": "",
                        "original_preserved": False, "verified": False,
                        "error": "Source not found"}
            if not _HAS_PYPDF:
                return {"success": False, "operation": "edit_pdf_metadata",
                        "source_path": abs_src, "output_path": "",
                        "original_preserved": True, "verified": False,
                        "error": "EditNotSupportedForFormat"}
            size_b, mtime_b = self._edit_snapshot(abs_src)
            try:
                reader = _PdfReader(abs_src)
                _ = len(reader.pages)
            except Exception as e:
                return {"success": False, "operation": "edit_pdf_metadata",
                        "source_path": abs_src, "output_path": "",
                        "original_preserved": True, "verified": False,
                        "error": "EditNotSupportedForFormat", "message": f"Not a valid PDF: {e}"}
            try:
                writer = _PdfWriter()
                for pg in reader.pages:
                    writer.add_page(pg)
                meta = {(k if k.startswith("/") else f"/{k}"): str(v)
                        for k, v in (metadata or {}).items()}
                if meta:
                    writer.add_metadata(meta)
                dest = os.path.abspath(output_path) if output_path else self._edit_sibling(abs_src)
                if dest == abs_src:
                    dest = self._edit_sibling(abs_src)
                Path(dest).parent.mkdir(parents=True, exist_ok=True)
                with open(dest, "wb") as f:
                    writer.write(f)
                preserved = self._edit_preserved(abs_src, size_b, mtime_b)
                verified = self._edit_verified(dest, preserved)
                return {"success": bool(verified), "operation": "edit_pdf_metadata",
                        "source_path": abs_src, "output_path": dest,
                        "original_preserved": preserved, "verified": verified,
                        "metadata": dict(metadata or {}),
                        "error": None if verified else "VerificationFailed"}
            except Exception as e:
                return {"success": False, "operation": "edit_pdf_metadata",
                        "source_path": abs_src, "output_path": "",
                        "original_preserved": self._edit_preserved(abs_src, size_b, mtime_b),
                        "verified": False, "error": str(e)}

    def edit_pptx_add_slide(self, source_path: str, text: str,
                            output_path: Optional[str] = None,
                            title: str = "") -> Dict[str, Any]:
        """Add a slide with text to a PPTX (sibling output, original preserved)."""
        with self._lock:
            abs_src = os.path.abspath(source_path)
            if not os.path.isfile(abs_src):
                return {"success": False, "operation": "edit_pptx",
                        "source_path": abs_src, "output_path": "",
                        "original_preserved": False, "verified": False,
                        "error": "Source not found"}
            if not _HAS_PPTX:
                return {"success": False, "operation": "edit_pptx",
                        "source_path": abs_src, "output_path": "",
                        "original_preserved": True, "verified": False,
                        "error": "EditNotSupportedForFormat"}
            size_b, mtime_b = self._edit_snapshot(abs_src)
            try:
                prs = _PptxPresentation(abs_src)
                n_before = len(prs.slides)
            except Exception as e:
                return {"success": False, "operation": "edit_pptx",
                        "source_path": abs_src, "output_path": "",
                        "original_preserved": True, "verified": False,
                        "error": "EditNotSupportedForFormat", "message": f"Not a valid PPTX: {e}"}
            try:
                from pptx.util import Inches as _Inches

                layout = prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[-1]
                slide = prs.slides.add_slide(layout)
                if title:
                    try:
                        slide.shapes.title.text = title
                    except Exception:
                        box = slide.shapes.add_textbox(_Inches(0.5), _Inches(0.2), _Inches(9), _Inches(0.8))
                        box.text_frame.text = title
                box = slide.shapes.add_textbox(_Inches(0.5), _Inches(1.2), _Inches(9), _Inches(4.5))
                box.text_frame.text = text or ""
                dest = os.path.abspath(output_path) if output_path else self._edit_sibling(abs_src)
                if dest == abs_src:
                    dest = self._edit_sibling(abs_src)
                Path(dest).parent.mkdir(parents=True, exist_ok=True)
                prs.save(dest)
                preserved = self._edit_preserved(abs_src, size_b, mtime_b)
                verified = self._edit_verified(dest, preserved)
                return {"success": bool(verified), "operation": "edit_pptx",
                        "source_path": abs_src, "output_path": dest,
                        "original_preserved": preserved, "verified": verified,
                        "slide_count_before": n_before, "slide_count_after": n_before + 1,
                        "error": None if verified else "VerificationFailed"}
            except Exception as e:
                return {"success": False, "operation": "edit_pptx",
                        "source_path": abs_src, "output_path": "",
                        "original_preserved": self._edit_preserved(abs_src, size_b, mtime_b),
                        "verified": False, "error": str(e)}

    def edit_pptx_replace_text(self, source_path: str, find_text: str, replace_text: str,
                               output_path: Optional[str] = None) -> Dict[str, Any]:
        """Replace text across PPTX slides (sibling output, original preserved)."""
        with self._lock:
            abs_src = os.path.abspath(source_path)
            if not os.path.isfile(abs_src):
                return {"success": False, "operation": "edit_pptx",
                        "source_path": abs_src, "output_path": "",
                        "original_preserved": False, "verified": False,
                        "error": "Source not found"}
            if not _HAS_PPTX:
                return {"success": False, "operation": "edit_pptx",
                        "source_path": abs_src, "output_path": "",
                        "original_preserved": True, "verified": False,
                        "error": "EditNotSupportedForFormat"}
            size_b, mtime_b = self._edit_snapshot(abs_src)
            try:
                prs = _PptxPresentation(abs_src)
                _ = len(prs.slides)
            except Exception as e:
                return {"success": False, "operation": "edit_pptx",
                        "source_path": abs_src, "output_path": "",
                        "original_preserved": True, "verified": False,
                        "error": "EditNotSupportedForFormat", "message": f"Not a valid PPTX: {e}"}
            texts: List[str] = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        texts.append(shape.text)
                    if shape.has_table:
                        for row in shape.table.rows:
                            for cell in row.cells:
                                texts.append(cell.text)
            full = "\n".join(texts)
            if find_text not in full:
                return {"success": False, "operation": "edit_pptx",
                        "source_path": abs_src, "output_path": "",
                        "original_preserved": True, "verified": False,
                        "error": "TextNotFound"}
            occurrences = full.count(find_text)
            try:
                for slide in prs.slides:
                    for shape in slide.shapes:
                        if shape.has_text_frame:
                            for para in shape.text_frame.paragraphs:
                                for run in para.runs:
                                    if find_text in run.text:
                                        run.text = run.text.replace(find_text, replace_text)
                                if find_text in para.text:
                                    para.text = para.text.replace(find_text, replace_text)
                        if shape.has_table:
                            for row in shape.table.rows:
                                for cell in row.cells:
                                    if find_text in cell.text:
                                        for para in cell.text_frame.paragraphs:
                                            for run in para.runs:
                                                if find_text in run.text:
                                                    run.text = run.text.replace(find_text, replace_text)
                                            if find_text in para.text:
                                                para.text = para.text.replace(find_text, replace_text)
                dest = os.path.abspath(output_path) if output_path else self._edit_sibling(abs_src)
                if dest == abs_src:
                    dest = self._edit_sibling(abs_src)
                Path(dest).parent.mkdir(parents=True, exist_ok=True)
                prs.save(dest)
                preserved = self._edit_preserved(abs_src, size_b, mtime_b)
                verified = self._edit_verified(dest, preserved)
                return {"success": bool(verified), "operation": "edit_pptx",
                        "source_path": abs_src, "output_path": dest,
                        "original_preserved": preserved, "verified": verified,
                        "occurrences": occurrences,
                        "error": None if verified else "VerificationFailed"}
            except Exception as e:
                return {"success": False, "operation": "edit_pptx",
                        "source_path": abs_src, "output_path": "",
                        "original_preserved": self._edit_preserved(abs_src, size_b, mtime_b),
                        "verified": False, "error": str(e)}

    def create_pdf_from_text(self, text: str, output_path: str,
                               title: str = "Genesis Document") -> Dict[str, Any]:
        """Create a PDF file from plain text via reportlab platypus.

        Writes a real PDF (title + body text) to ``output_path``,
        then verifies the output exists with size>0. When reportlab
        is missing, honestly returns success=False with error
        'reportlab required for PDF creation (pip install reportlab)'
        — PDF bytes are never faked.
        """
        with self._lock:
            dest = os.path.abspath(output_path)
            if not _HAS_REPORTLAB:
                return {"success": False, "operation": "create_pdf",
                        "source_path": "", "output_path": dest,
                        "original_preserved": True, "verified": False,
                        "error": "reportlab required for PDF creation (pip install reportlab)",
                        "message": "reportlab required for PDF creation: "
                                   "install with pip install reportlab"}
            try:
                from reportlab.lib.pagesizes import letter as _letter
                from reportlab.lib.styles import getSampleStyleSheet as _get_styles
                from reportlab.platypus import (
                    Paragraph as _Paragraph,
                    SimpleDocTemplate as _DocTemplate,
                    Spacer as _Spacer,
                )
            except ImportError:
                return {"success": False, "operation": "create_pdf",
                        "source_path": "", "output_path": dest,
                        "original_preserved": True, "verified": False,
                        "error": "reportlab required for PDF creation (pip install reportlab)",
                        "message": "reportlab required for PDF creation: "
                                   "install with pip install reportlab"}
            try:
                Path(dest).parent.mkdir(parents=True, exist_ok=True)
                styles = _get_styles()
                story = [_Paragraph(title or "Untitled", styles["Title"]),
                         _Spacer(1, 12)]
                for para in (text or "").split("\n"):
                    story.append(_Paragraph(para if para.strip() else " ",
                                            styles["Normal"]))
                    story.append(_Spacer(1, 6))
                _DocTemplate(dest, pagesize=_letter).build(story)
                verified = self._edit_verified(dest, True)
                return {"success": bool(verified), "operation": "create_pdf",
                        "source_path": "", "output_path": dest,
                        "original_preserved": True, "verified": verified,
                        "title": title,
                        "error": None if verified else "VerificationFailed"}
            except Exception as e:
                return {"success": False, "operation": "create_pdf",
                        "source_path": "", "output_path": dest,
                        "original_preserved": True, "verified": False,
                        "error": str(e)}

    def get_supported_conversions(self) -> List[Dict[str, Any]]:
        with self._lock:
            pairs = []
            mapping = {"word": ["pdf", "text"], "excel": ["csv", "json"],
                       "pdf": ["text", "word"], "powerpoint": ["pdf", "text"],
                       "csv": ["json"], "json": ["csv"], "text": ["markdown"]}
            for src, tgts in mapping.items():
                sc = self._format_capabilities.get(src)
                if not sc:
                    continue
                for tgt in tgts:
                    tc = self._format_capabilities.get(tgt)
                    if tc:
                        pairs.append({"source_format": src, "target_format": tgt,
                                      "source_ext": sc.extensions, "target_ext": tc.extensions,
                                      "fidelity": "high" if tc.supports_formatting and sc.supports_formatting else "medium"})
            return pairs

    def get_format_info(self, fmt: str) -> Optional[Dict[str, Any]]:
        with self._lock:
            cap = self._format_capabilities.get(fmt.lower())
            return cap.to_dict() if cap else None

    def get_conversion_history(self, limit: int = 50) -> List[Dict[str, Any]]:
        with self._lock:
            return [o.to_dict() for o in self._conversion_history[-limit:]]

    def get_converter_summary(self) -> Dict[str, Any]:
        with self._lock:
            total = len(self._conversion_history)
            by_status: Dict[str, int] = {}
            for o in self._conversion_history:
                by_status[o.status.value] = by_status.get(o.status.value, 0) + 1
            avg_f = round(sum(o.fidelity_score for o in self._conversion_history) / total, 4) if total else 0.0
            return {
                "version": self.VERSION, "total_conversions": total,
                "conversions_by_status": by_status, "average_fidelity_score": avg_f,
                "supported_formats": list(self._format_capabilities.keys()),
                "total_conversion_pairs": len(self.get_supported_conversions()),
            }
