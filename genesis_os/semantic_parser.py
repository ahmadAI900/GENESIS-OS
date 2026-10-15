from __future__ import annotations
import os, hashlib, json, time, threading, re, pathlib, dataclasses, enum
from typing import Any, Dict, List, Optional, Set
from dataclasses import dataclass, field

DOCX_AVAILABLE = OPENPYXL_AVAILABLE = PPTX_AVAILABLE = PYPDF_AVAILABLE = False
try:
    import docx; DOCX_AVAILABLE = True
except ImportError: pass
try:
    import openpyxl; OPENPYXL_AVAILABLE = True
except ImportError: pass
try:
    from pptx import Presentation; PPTX_AVAILABLE = True
except ImportError: pass
try:
    from pypdf import PdfReader; PYPDF_AVAILABLE = True
except ImportError: pass


class DocumentType(str, enum.Enum):
    WORD="word"; EXCEL="excel"; POWERPOINT="powerpoint"; PDF="pdf"; TEXT="text"
    CSV="csv"; JSON="json"; XML="xml"; MARKDOWN="markdown"; UNKNOWN="unknown"

class ElementType(str, enum.Enum):
    HEADING="heading"; PARAGRAPH="paragraph"; TABLE="table"; LIST="list"
    IMAGE="image"; CODE_BLOCK="code_block"; QUOTE="quote"; LINK="link"
    REFERENCE="reference"; FOOTNOTE="footnote"; METADATA="metadata"


@dataclass
class DocumentElement:
    element_id: str; element_type: ElementType; content: str; level: int = 0
    metadata: Dict[str, Any] = field(default_factory=dict)
    children: List[DocumentElement] = field(default_factory=list)
    references: List[str] = field(default_factory=list)
    def to_dict(self) -> Dict[str, Any]:
        return {"element_id": self.element_id, "element_type": self.element_type.value,
                "content": self.content, "level": self.level, "metadata": self.metadata,
                "children": [c.to_dict() for c in self.children], "references": self.references}


@dataclass
class DocumentStructure:
    doc_id: str; file_path: str; doc_type: DocumentType; title: str; author: str
    created_at: str; modified_at: str; language: str; word_count: int; page_count: int
    elements: List[DocumentElement]; metadata: Dict[str, Any]; semantic_hash: str
    sections: List[Dict[str, Any]]; tables: List[Dict[str, Any]]
    images_count: int; links: List[str]
    def to_dict(self) -> Dict[str, Any]:
        return {"doc_id": self.doc_id, "file_path": self.file_path,
                "doc_type": self.doc_type.value, "title": self.title, "author": self.author,
                "created_at": self.created_at, "modified_at": self.modified_at,
                "language": self.language, "word_count": self.word_count,
                "page_count": self.page_count, "elements": [e.to_dict() for e in self.elements],
                "metadata": self.metadata, "semantic_hash": self.semantic_hash,
                "sections": self.sections, "tables": self.tables,
                "images_count": self.images_count, "links": self.links}


@dataclass
class DocumentRelationship:
    relationship_id: str; source_doc_id: str; target_doc_id: str
    relationship_type: str; confidence: float; evidence: str
    def to_dict(self) -> Dict[str, Any]:
        return {"relationship_id": self.relationship_id, "source_doc_id": self.source_doc_id,
                "target_doc_id": self.target_doc_id, "relationship_type": self.relationship_type,
                "confidence": self.confidence, "evidence": self.evidence}


class SemanticDocumentParser:
    VERSION = "1.0.0"

    def __init__(self, max_docs: int = 500):
        self._parsed_docs: Dict[str, DocumentStructure] = {}
        self._relationships: List[DocumentRelationship] = []
        self._doc_graph: Dict[str, Set[str]] = {}
        self._lock = threading.RLock()
        self._max_docs = max_docs

    def _detect_doc_type(self, file_path: str) -> DocumentType:
        ext = pathlib.Path(file_path).suffix.lower()
        ext_map = {".docx": DocumentType.WORD, ".doc": DocumentType.WORD,
                   ".xlsx": DocumentType.EXCEL, ".xls": DocumentType.EXCEL,
                   ".pptx": DocumentType.POWERPOINT, ".ppt": DocumentType.POWERPOINT,
                   ".pdf": DocumentType.PDF, ".txt": DocumentType.TEXT,
                   ".csv": DocumentType.CSV, ".json": DocumentType.JSON,
                   ".xml": DocumentType.XML, ".md": DocumentType.MARKDOWN,
                   ".markdown": DocumentType.MARKDOWN}
        if ext in ext_map: return ext_map[ext]
        try:
            with open(file_path, "rb") as f: header = f.read(8)
            if header[:4] == b"%PDF": return DocumentType.PDF
            if header[:4] in (b"PK\x03\x04", b"\xd0\xcf\x11\xe0"): return DocumentType.WORD
        except Exception: pass
        return DocumentType.UNKNOWN

    def _compute_semantic_hash(self, content: bytes) -> str:
        return hashlib.sha256(content).hexdigest()

    def parse(self, file_path: str) -> DocumentStructure:
        with self._lock:
            doc_type = self._detect_doc_type(file_path)
            parsers = {DocumentType.WORD: self._parse_word, DocumentType.EXCEL: self._parse_excel,
                       DocumentType.POWERPOINT: self._parse_powerpoint, DocumentType.PDF: self._parse_pdf,
                       DocumentType.TEXT: self._parse_text, DocumentType.CSV: self._parse_csv,
                       DocumentType.JSON: self._parse_json, DocumentType.MARKDOWN: self._parse_text,
                       DocumentType.XML: self._parse_text}
            structure = parsers.get(doc_type, self._parse_text)(file_path)
            if len(self._parsed_docs) >= self._max_docs:
                self._parsed_docs.pop(next(iter(self._parsed_docs)))
            self._parsed_docs[structure.doc_id] = structure
            self._doc_graph.setdefault(structure.doc_id, set())
            return structure

    def _parse_word(self, file_path: str) -> DocumentStructure:
        try:
            if not DOCX_AVAILABLE: return self._fallback(file_path, DocumentType.WORD)
            doc = docx.Document(file_path)
            elements, tables, links = [], self._extract_tables_from_word(doc), []
            for i, para in enumerate(doc.paragraphs):
                style = para.style.name.lower() if para.style else ""
                if "heading" in style:
                    m = re.search(r"(\d)", style)
                    el_type, level = ElementType.HEADING, int(m.group(1)) if m else 1
                else:
                    el_type, level = ElementType.PARAGRAPH, 0
                elements.append(DocumentElement(f"word_el_{i}", el_type, para.text, level, {"style": style}))
            word_count = sum(len(e.content.split()) for e in elements)
            core = doc.core_properties
            return DocumentStructure(hashlib.md5(file_path.encode()).hexdigest()[:16], file_path,
                DocumentType.WORD, core.title or pathlib.Path(file_path).stem, core.author or "",
                str(core.created) if core.created else "", str(core.modified) if core.modified else "",
                core.language or "en", word_count, 1, elements, {"tables_count": len(tables)},
                self._compute_semantic_hash(open(file_path, "rb").read()),
                self._extract_sections(elements), tables, 0, links)
        except Exception: return self._fallback(file_path, DocumentType.WORD)

    def _parse_excel(self, file_path: str) -> DocumentStructure:
        try:
            if not OPENPYXL_AVAILABLE: return self._fallback(file_path, DocumentType.EXCEL)
            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            elements, tables = [], []
            for name in wb.sheetnames:
                ws = wb[name]
                for t in self._extract_tables_from_excel(ws):
                    t["sheet"] = name; tables.append(t)
                rows = list(ws.iter_rows(values_only=True))
                if rows:
                    headers = [str(c) if c else f"col_{i}" for i, c in enumerate(rows[0])]
                    elements.append(DocumentElement(f"sheet_{name}", ElementType.TABLE,
                        json.dumps(headers), 0, {"sheet_name": name, "rows": len(rows)}))
            word_count = sum(len(e.content.split()) for e in elements)
            return DocumentStructure(hashlib.md5(file_path.encode()).hexdigest()[:16], file_path,
                DocumentType.EXCEL, pathlib.Path(file_path).stem, "", "", "", "en",
                word_count, 0, elements, {"sheet_count": len(wb.sheetnames), "sheets": wb.sheetnames},
                self._compute_semantic_hash(open(file_path, "rb").read()), [], tables, 0, [])
        except Exception: return self._fallback(file_path, DocumentType.EXCEL)

    def _parse_powerpoint(self, file_path: str) -> DocumentStructure:
        try:
            if not PPTX_AVAILABLE: return self._fallback(file_path, DocumentType.POWERPOINT)
            prs = Presentation(file_path)
            elements, img_count, tables = [], 0, []
            for si, slide in enumerate(prs.slides):
                texts = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip(): texts.append(shape.text)
                    if shape.shape_type == 13: img_count += 1
                    if shape.has_table:
                        rows = [[c.text for c in row.cells] for row in shape.table.rows]
                        if rows: tables.append({"headers": rows[0], "data": rows[1:], "slide": si})
                elements.append(DocumentElement(f"slide_{si}", ElementType.PARAGRAPH,
                    "\n".join(texts), 0, {"slide_number": si + 1}))
            word_count = sum(len(e.content.split()) for e in elements)
            return DocumentStructure(hashlib.md5(file_path.encode()).hexdigest()[:16], file_path,
                DocumentType.POWERPOINT, pathlib.Path(file_path).stem, "", "", "", "en",
                word_count, len(prs.slides), elements, {"slide_count": len(prs.slides)},
                self._compute_semantic_hash(open(file_path, "rb").read()), [], tables, img_count, [])
        except Exception: return self._fallback(file_path, DocumentType.POWERPOINT)

    def _parse_pdf(self, file_path: str) -> DocumentStructure:
        try:
            if not PYPDF_AVAILABLE: return self._fallback(file_path, DocumentType.PDF)
            reader = PdfReader(file_path)
            elements = []
            for pi, page in enumerate(reader.pages):
                elements.append(DocumentElement(f"pdf_page_{pi}", ElementType.PARAGRAPH,
                    page.extract_text() or "", 0, {"page_number": pi + 1}))
            info = reader.metadata or {}
            links = []
            for page in reader.pages:
                if "/Annots" in page:
                    for annot in page["/Annots"]:
                        obj = annot.get_object()
                        if "/A" in obj and "/URI" in obj["/A"]: links.append(str(obj["/A"]["/URI"]))
            word_count = sum(len(e.content.split()) for e in elements)
            return DocumentStructure(hashlib.md5(file_path.encode()).hexdigest()[:16], file_path,
                DocumentType.PDF, str(info.get("/Title", pathlib.Path(file_path).stem)),
                str(info.get("/Author", "")), str(info.get("/CreationDate", "")),
                str(info.get("/ModDate", "")), str(info.get("/Language", "en")),
                word_count, len(reader.pages), elements, {"producer": str(info.get("/Producer", ""))},
                self._compute_semantic_hash(open(file_path, "rb").read()),
                self._extract_sections(elements), [], 0, links)
        except Exception: return self._fallback(file_path, DocumentType.PDF)

    def _parse_text(self, file_path: str) -> DocumentStructure:
        try:
            content = open(file_path, "r", encoding="utf-8", errors="replace").read()
            elements, idx, in_code, code_buf = [], 0, False, []
            for line in content.split("\n"):
                if line.strip().startswith("```"):
                    if in_code:
                        elements.append(DocumentElement(f"txt_{idx}", ElementType.CODE_BLOCK, "\n".join(code_buf)))
                        idx += 1; code_buf = []; in_code = False
                    else: in_code = True
                    continue
                if in_code: code_buf.append(line); continue
                hm = re.match(r"^(#{1,6})\s+(.+)", line)
                if hm:
                    elements.append(DocumentElement(f"txt_{idx}", ElementType.HEADING, hm.group(2).strip(), len(hm.group(1))))
                elif line.strip().startswith("> "):
                    elements.append(DocumentElement(f"txt_{idx}", ElementType.QUOTE, line.strip()[2:]))
                elif re.match(r"^[-*]\s+", line.strip()):
                    elements.append(DocumentElement(f"txt_{idx}", ElementType.LIST, line.strip()[2:], 1))
                elif line.strip():
                    elements.append(DocumentElement(f"txt_{idx}", ElementType.PARAGRAPH, line.strip()))
                else: continue
                idx += 1
            word_count = sum(len(e.content.split()) for e in elements)
            return DocumentStructure(hashlib.md5(file_path.encode()).hexdigest()[:16], file_path,
                self._detect_doc_type(file_path), pathlib.Path(file_path).stem, "", "", "", "en",
                word_count, 1, elements, {}, self._compute_semantic_hash(content.encode()),
                self._extract_sections(elements), [], 0, [])
        except Exception: return self._fallback(file_path, DocumentType.TEXT)

    def _parse_csv(self, file_path: str) -> DocumentStructure:
        try:
            content = open(file_path, "r", encoding="utf-8", errors="replace").read()
            lines = content.strip().split("\n")
            if not lines: return self._fallback(file_path, DocumentType.CSV)
            delim = next((d for d in [",", ";", "\t", "|"] if d in lines[0]), ",")
            headers = [h.strip() for h in lines[0].split(delim)]
            sample = lines[1].split(delim) if len(lines) > 1 else []
            col_types = ["integer" if re.match(r"^-?\d+$", s.strip()) else
                         "float" if re.match(r"^-?\d+\.\d+$", s.strip()) else "string"
                         for s in sample[:len(headers)]]
            while len(col_types) < len(headers): col_types.append("string")
            table = {"headers": headers, "data_rows": len(lines) - 1, "column_types": col_types, "delimiter": delim}
            el = DocumentElement("csv_table", ElementType.TABLE, json.dumps(headers),
                                 0, {"row_count": len(lines) - 1, "column_count": len(headers)})
            return DocumentStructure(hashlib.md5(file_path.encode()).hexdigest()[:16], file_path,
                DocumentType.CSV, pathlib.Path(file_path).stem, "", "", "", "en",
                len(lines), 1, [el], {}, self._compute_semantic_hash(content.encode()),
                [], [table], 0, [])
        except Exception: return self._fallback(file_path, DocumentType.CSV)

    def _parse_json(self, file_path: str) -> DocumentStructure:
        try:
            content = open(file_path, "r", encoding="utf-8", errors="replace").read()
            data = json.loads(content)
            schema = self._analyze_json(data)
            el = DocumentElement("json_root", ElementType.METADATA, json.dumps(schema, indent=2)[:500], 0, schema)
            return DocumentStructure(hashlib.md5(file_path.encode()).hexdigest()[:16], file_path,
                DocumentType.JSON, pathlib.Path(file_path).stem, "", "", "", "en",
                0, 1, [el], schema, self._compute_semantic_hash(content.encode()), [], [], 0, [])
        except Exception: return self._fallback(file_path, DocumentType.JSON)

    def _analyze_json(self, data: Any, depth: int = 0) -> Dict[str, Any]:
        if depth > 10: return {"type": "max_depth_exceeded"}
        if isinstance(data, dict):
            return {"type": "object", "keys": list(data.keys())[:20], "key_count": len(data),
                    "nesting_depth": depth, "value_types": {k: type(v).__name__ for k, v in list(data.items())[:10]}}
        elif isinstance(data, list):
            return {"type": "array", "length": len(data), "item_types": [type(i).__name__ for i in data[:5]]}
        elif isinstance(data, str): return {"type": "string", "length": len(data)}
        elif isinstance(data, (int, float)): return {"type": "numeric", "value": data}
        elif isinstance(data, bool): return {"type": "boolean"}
        elif data is None: return {"type": "null"}
        return {"type": str(type(data).__name__)}

    def _extract_sections(self, elements: List[DocumentElement]) -> List[Dict[str, Any]]:
        sections, stack = [], []
        for el in elements:
            if el.element_type == ElementType.HEADING and el.level > 0:
                sec = {"id": el.element_id, "title": el.content, "level": el.level, "children": []}
                while stack and stack[-1]["level"] >= el.level: stack.pop()
                if stack: stack[-1]["children"].append(sec)
                else: sections.append(sec)
                stack.append(sec)
        return sections

    def _extract_tables_from_word(self, doc) -> List[Dict[str, Any]]:
        tables = []
        for ti, table in enumerate(doc.tables):
            rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            if rows: tables.append({"id": f"word_table_{ti}", "headers": rows[0],
                                     "data": rows[1:] if len(rows) > 1 else [], "row_count": len(rows),
                                     "col_count": len(rows[0]) if rows else 0})
        return tables

    def _extract_tables_from_excel(self, ws) -> List[Dict[str, Any]]:
        rows = list(ws.iter_rows(values_only=True))
        if not rows: return []
        headers = [str(c) if c else f"col_{i}" for i, c in enumerate(rows[0])]
        data = [[str(c) if c is not None else "" for c in row] for row in rows[1:]]
        col_types = []
        for ci in range(len(headers)):
            vals = [r[ci] for r in data if ci < len(r)]
            if vals and all(re.match(r"^-?\d+$", str(v)) for v in vals[:5]): col_types.append("integer")
            elif vals and all(re.match(r"^-?\d+\.\d+$", str(v)) for v in vals[:5]): col_types.append("float")
            else: col_types.append("string")
        return [{"headers": headers, "data": data[:100], "row_count": len(data),
                 "col_count": len(headers), "column_types": col_types}]

    def detect_relationships(self, doc1: DocumentStructure, doc2: DocumentStructure) -> Optional[DocumentRelationship]:
        with self._lock:
            if doc1.doc_id == doc2.doc_id: return None
            evidence, confidence, rel_type = [], 0.0, "similar_to"
            if doc1.title and doc2.title:
                overlap = len(set(doc1.title.lower().split()) & set(doc2.title.lower().split()))
                if overlap > 0: confidence += min(overlap * 0.15, 0.3); evidence.append(f"Title overlap: {overlap}")
            c1 = set(" ".join(e.content for e in doc1.elements[:20]).lower().split())
            c2 = set(" ".join(e.content for e in doc2.elements[:20]).lower().split())
            if c1 and c2:
                ov = len(c1 & c2) / max(len(c1 | c2), 1)
                confidence += ov * 0.4; evidence.append(f"Content overlap: {ov:.2%}")
            d2_lower = doc2.title.lower()
            if d2_lower in " ".join(e.content.lower() for e in doc1.elements[:30]):
                confidence += 0.3; rel_type = "references"; evidence.append(f"{doc2.title} referenced")
            d1_lower = doc1.title.lower()
            if d1_lower in " ".join(e.content.lower() for e in doc2.elements[:30]):
                confidence += 0.3; rel_type = "references"; evidence.append(f"{doc1.title} referenced")
            if confidence < 0.1: return None
            rel = DocumentRelationship(f"rel_{doc1.doc_id[:8]}_{doc2.doc_id[:8]}",
                doc1.doc_id, doc2.doc_id, rel_type, min(round(confidence, 3), 1.0), "; ".join(evidence))
            self._relationships.append(rel)
            self._doc_graph.setdefault(doc1.doc_id, set()).add(doc2.doc_id)
            self._doc_graph.setdefault(doc2.doc_id, set()).add(doc1.doc_id)
            return rel

    def build_document_graph(self) -> Dict[str, Any]:
        with self._lock:
            docs = list(self._parsed_docs.values())
            for i in range(len(docs)):
                for j in range(i + 1, len(docs)):
                    exists = any({r.source_doc_id, r.target_doc_id} == {docs[i].doc_id, docs[j].doc_id}
                                 for r in self._relationships)
                    if not exists: self.detect_relationships(docs[i], docs[j])
            return {"documents": len(self._parsed_docs), "relationships": len(self._relationships),
                    "graph": {k: list(v) for k, v in self._doc_graph.items()}}

    def search_documents(self, query: str, limit: int = 10) -> List[Dict[str, Any]]:
        with self._lock:
            qw = set(query.lower().split())
            results = []
            for did, doc in self._parsed_docs.items():
                score = (1.0 if query.lower() in doc.title.lower() else 0.0)
                cw = set(" ".join(e.content for e in doc.elements).lower().split())
                if qw: score += len(qw & cw) / max(len(qw), 1)
                if score > 0:
                    matched = [{"id": e.element_id, "type": e.element_type.value, "content": e.content[:200]}
                               for e in doc.elements if any(w in e.content.lower() for w in qw)][:5]
                    results.append({"doc_id": did, "title": doc.title, "doc_type": doc.doc_type.value,
                                    "score": round(score, 3), "matched_elements": matched})
            results.sort(key=lambda x: x["score"], reverse=True)
            return results[:limit]

    def get_document_summary(self, doc_id: str) -> Dict[str, Any]:
        with self._lock:
            if doc_id not in self._parsed_docs: return {"error": f"Document {doc_id} not found"}
            doc = self._parsed_docs[doc_id]
            rels = [r.to_dict() for r in self._relationships
                    if r.source_doc_id == doc_id or r.target_doc_id == doc_id]
            return {"doc_id": doc_id, "title": doc.title, "doc_type": doc.doc_type.value,
                    "word_count": doc.word_count, "page_count": doc.page_count,
                    "element_count": len(doc.elements), "section_count": len(doc.sections),
                    "table_count": len(doc.tables), "images_count": doc.images_count,
                    "link_count": len(doc.links), "relationships": rels,
                    "related_docs": list(self._doc_graph.get(doc_id, set()))}

    def get_parser_summary(self) -> Dict[str, Any]:
        with self._lock:
            type_counts = {}
            for doc in self._parsed_docs.values():
                t = doc.doc_type.value; type_counts[t] = type_counts.get(t, 0) + 1
            total_el = sum(len(d.elements) for d in self._parsed_docs.values())
            return {"version": self.VERSION, "total_documents": len(self._parsed_docs),
                    "total_relationships": len(self._relationships), "total_elements": total_el,
                    "documents_by_type": type_counts, "graph_nodes": len(self._doc_graph),
                    "graph_edges": sum(len(v) for v in self._doc_graph.values()) // 2,
                    "max_docs": self._max_docs,
                    "available_parsers": {"word": DOCX_AVAILABLE, "excel": OPENPYXL_AVAILABLE,
                                          "powerpoint": PPTX_AVAILABLE, "pdf": PYPDF_AVAILABLE}}

    def _fallback(self, file_path: str, doc_type: DocumentType) -> DocumentStructure:
        try: content = open(file_path, "r", encoding="utf-8", errors="replace").read()
        except Exception: content = ""
        el = DocumentElement("fallback_0", ElementType.PARAGRAPH, content[:10000])
        return DocumentStructure(hashlib.md5(file_path.encode()).hexdigest()[:16], file_path, doc_type,
            pathlib.Path(file_path).stem, "", "", "", "en", len(content.split()), 1, [el],
            {"fallback": True}, self._compute_semantic_hash(content.encode()), [], [], 0, [])
