from __future__ import annotations

import os
import json
import time
import threading
import re
import hashlib
import statistics
from dataclasses import dataclass, field, asdict
from enum import Enum
from typing import Any, Dict, List, Optional, Tuple

try:
    import openpyxl
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

try:
    import csv
    HAS_CSV = True
except ImportError:
    HAS_CSV = False


class DataType(str, Enum):
    STRING = "string"
    INTEGER = "integer"
    FLOAT = "float"
    BOOLEAN = "boolean"
    DATE = "date"
    DATETIME = "datetime"
    CURRENCY = "currency"
    PERCENTAGE = "percentage"
    EMAIL = "email"
    URL = "url"
    PHONE = "phone"
    ADDRESS = "address"
    UNKNOWN = "unknown"


@dataclass
class ColumnProfile:
    column_name: str
    data_type: DataType
    nullable: bool
    unique_count: int
    total_count: int
    null_count: int
    min_value: Optional[str]
    max_value: Optional[str]
    mean_value: Optional[float]
    std_dev: Optional[float]
    sample_values: List[str]
    format_pattern: Optional[str] = None

    def to_dict(self) -> Dict[str, Any]:
        return {
            "column_name": self.column_name,
            "data_type": self.data_type.value,
            "nullable": self.nullable,
            "unique_count": self.unique_count,
            "total_count": self.total_count,
            "null_count": self.null_count,
            "min_value": self.min_value,
            "max_value": self.max_value,
            "mean_value": self.mean_value,
            "std_dev": self.std_dev,
            "sample_values": self.sample_values,
            "format_pattern": self.format_pattern,
        }


@dataclass
class DataQualityIssue:
    issue_id: str
    issue_type: str
    column: Optional[str]
    row: Optional[int]
    description: str
    severity: str
    suggestion: str
    auto_fixable: bool

    def to_dict(self) -> Dict[str, Any]:
        return {
            "issue_id": self.issue_id,
            "issue_type": self.issue_type,
            "column": self.column,
            "row": self.row,
            "description": self.description,
            "severity": self.severity,
            "suggestion": self.suggestion,
            "auto_fixable": self.auto_fixable,
        }


@dataclass
class SchemaProfile:
    schema_id: str
    file_path: str
    file_type: str
    sheet_name: Optional[str]
    row_count: int
    column_count: int
    columns: List[ColumnProfile]
    primary_key_candidates: List[str]
    foreign_key_candidates: List[Dict[str, str]]
    data_quality_score: float
    quality_issues: List[DataQualityIssue]
    relationships: List[Dict[str, Any]]
    metadata: Dict[str, Any]

    def to_dict(self) -> Dict[str, Any]:
        return {
            "schema_id": self.schema_id,
            "file_path": self.file_path,
            "file_type": self.file_type,
            "sheet_name": self.sheet_name,
            "row_count": self.row_count,
            "column_count": self.column_count,
            "columns": [c.to_dict() for c in self.columns],
            "primary_key_candidates": self.primary_key_candidates,
            "foreign_key_candidates": self.foreign_key_candidates,
            "data_quality_score": self.data_quality_score,
            "quality_issues": [q.to_dict() for q in self.quality_issues],
            "relationships": self.relationships,
            "metadata": self.metadata,
        }


class SchemaEngine:
    VERSION = "1.0.0"

    def __init__(self) -> None:
        self._profiles: Dict[str, SchemaProfile] = {}
        self._lock = threading.RLock()

    def profile_file(self, file_path: str, sheet_name: Optional[str] = None) -> SchemaProfile:
        with self._lock:
            ext = os.path.splitext(file_path)[1].lower()
            if ext in (".xlsx", ".xls"):
                profile = self._profile_excel(file_path, sheet_name)
            elif ext == ".csv":
                profile = self._profile_csv(file_path)
            elif ext == ".json":
                profile = self._profile_json(file_path)
            else:
                raise ValueError(f"Unsupported file type: {ext}")
            self._profiles[profile.schema_id] = profile
            return profile

    def _profile_excel(self, file_path: str, sheet_name: Optional[str] = None) -> SchemaProfile:
        if not HAS_OPENPYXL:
            raise ImportError("openpyxl required for Excel profiling")
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        target_sheet = sheet_name or wb.sheetnames[0]
        ws = wb[target_sheet]
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
        if not rows:
            return self._empty_profile(file_path, "excel", target_sheet)
        headers = [str(h) if h is not None else f"col_{i}" for i, h in enumerate(rows[0])]
        data_rows = rows[1:]
        columns = self._build_column_profiles(headers, data_rows)
        issues = self._detect_quality_issues(columns, len(data_rows))
        score = self._compute_quality_score(columns, issues)
        pk = self._detect_primary_key(columns)
        return SchemaProfile(
            schema_id=hashlib.sha256(f"{file_path}:{target_sheet}".encode()).hexdigest()[:16],
            file_path=file_path,
            file_type="excel",
            sheet_name=target_sheet,
            row_count=len(data_rows),
            column_count=len(headers),
            columns=columns,
            primary_key_candidates=pk,
            foreign_key_candidates=[],
            data_quality_score=score,
            quality_issues=issues,
            relationships=[],
            metadata={"engine_version": self.VERSION, "profiled_at": time.time()},
        )

    def _profile_csv(self, file_path: str) -> SchemaProfile:
        if not HAS_CSV:
            raise ImportError("csv module required")
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            reader = csv.reader(f)
            rows = list(reader)
        if not rows:
            return self._empty_profile(file_path, "csv", None)
        headers = rows[0]
        data_rows = [r for r in rows[1:] if any(cell.strip() for cell in r)]
        columns = self._build_column_profiles(headers, data_rows)
        issues = self._detect_quality_issues(columns, len(data_rows))
        score = self._compute_quality_score(columns, issues)
        pk = self._detect_primary_key(columns)
        return SchemaProfile(
            schema_id=hashlib.sha256(file_path.encode()).hexdigest()[:16],
            file_path=file_path,
            file_type="csv",
            sheet_name=None,
            row_count=len(data_rows),
            column_count=len(headers),
            columns=columns,
            primary_key_candidates=pk,
            foreign_key_candidates=[],
            data_quality_score=score,
            quality_issues=issues,
            relationships=[],
            metadata={"engine_version": self.VERSION, "profiled_at": time.time()},
        )

    def _profile_json(self, file_path: str) -> SchemaProfile:
        with open(file_path, "r", encoding="utf-8") as f:
            data = json.load(f)
        records = data if isinstance(data, list) else [data]
        if not records:
            return self._empty_profile(file_path, "json", None)
        all_keys: List[str] = []
        for rec in records:
            for k in rec.keys():
                if k not in all_keys:
                    all_keys.append(k)
        data_rows = [[rec.get(k) for k in all_keys] for rec in records]
        columns = self._build_column_profiles(all_keys, data_rows)
        issues = self._detect_quality_issues(columns, len(data_rows))
        score = self._compute_quality_score(columns, issues)
        pk = self._detect_primary_key(columns)
        return SchemaProfile(
            schema_id=hashlib.sha256(file_path.encode()).hexdigest()[:16],
            file_path=file_path,
            file_type="json",
            sheet_name=None,
            row_count=len(data_rows),
            column_count=len(all_keys),
            columns=columns,
            primary_key_candidates=pk,
            foreign_key_candidates=[],
            data_quality_score=score,
            quality_issues=issues,
            relationships=[],
            metadata={"engine_version": self.VERSION, "profiled_at": time.time()},
        )

    def _empty_profile(self, file_path: str, file_type: str, sheet: Optional[str]) -> SchemaProfile:
        return SchemaProfile(
            schema_id=hashlib.sha256(f"{file_path}:empty".encode()).hexdigest()[:16],
            file_path=file_path, file_type=file_type, sheet_name=sheet,
            row_count=0, column_count=0, columns=[], primary_key_candidates=[],
            foreign_key_candidates=[], data_quality_score=100.0, quality_issues=[],
            relationships=[], metadata={"engine_version": self.VERSION},
        )

    def _build_column_profiles(self, headers: List[str], rows: List[Any]) -> List[ColumnProfile]:
        columns: List[ColumnProfile] = []
        for ci, header in enumerate(headers):
            values = [row[ci] if ci < len(row) else None for row in rows]
            non_null = [v for v in values if v is not None and str(v).strip() != ""]
            null_count = len(values) - len(non_null)
            str_vals = [str(v) for v in non_null]
            unique = len(set(str_vals))
            dtype = self._detect_column_type(non_null)
            fmt = self._detect_format_pattern(str_vals, dtype) if str_vals else None
            mn = min(str_vals) if str_vals else None
            mx = max(str_vals) if str_vals else None
            mean_v = None
            std_v = None
            if dtype in (DataType.INTEGER, DataType.FLOAT, DataType.CURRENCY, DataType.PERCENTAGE):
                nums = []
                for s in str_vals:
                    try:
                        nums.append(float(re.sub(r"[,$%]", "", s)))
                    except ValueError:
                        pass
                if nums:
                    mean_v = round(statistics.mean(nums), 4)
                    std_v = round(statistics.stdev(nums), 4) if len(nums) > 1 else 0.0
            samples = str_vals[:5]
            columns.append(ColumnProfile(
                column_name=header, data_type=dtype, nullable=null_count > 0,
                unique_count=unique, total_count=len(values), null_count=null_count,
                min_value=mn, max_value=mx, mean_value=mean_v, std_dev=std_v,
                sample_values=samples, format_pattern=fmt,
            ))
        return columns

    def _detect_column_type(self, values: List[Any]) -> DataType:
        if not values:
            return DataType.UNKNOWN
        samples = [str(v).strip() for v in values[:200]]
        bool_true = {"true", "yes", "1"}
        bool_false = {"false", "no", "0"}
        if all(s.lower() in bool_true | bool_false for s in samples):
            return DataType.BOOLEAN
        email_re = re.compile(r"^[\w.\-+]+@[\w.\-]+\.\w+$")
        if sum(1 for s in samples if email_re.match(s)) / max(len(samples), 1) > 0.7:
            return DataType.EMAIL
        url_re = re.compile(r"^https?://\S+$")
        if sum(1 for s in samples if url_re.match(s)) / max(len(samples), 1) > 0.7:
            return DataType.URL
        phone_re = re.compile(r"^[\+]?[\d\s\-\(\)]{7,}$")
        if sum(1 for s in samples if phone_re.match(s)) / max(len(samples), 1) > 0.7:
            return DataType.PHONE
        pct_re = re.compile(r"^-?\d+\.?\d*\s*%$")
        if sum(1 for s in samples if pct_re.match(s)) / max(len(samples), 1) > 0.7:
            return DataType.PERCENTAGE
        curr_re = re.compile(r"^[$€£]\s?[\d,]+\.?\d*$")
        if sum(1 for s in samples if curr_re.match(s)) / max(len(samples), 1) > 0.7:
            return DataType.CURRENCY
        date_formats = ["%Y-%m-%d", "%m/%d/%Y", "%d/%m/%Y", "%Y/%m/%d", "%Y-%m-%dT%H:%M:%S"]
        for fmt in date_formats:
            try:
                parsed = [time.strptime(s, fmt) for s in samples[:10]]
                if len(parsed) >= len(samples[:10]) * 0.7:
                    if "T" in fmt or "%H" in fmt:
                        return DataType.DATETIME
                    return DataType.DATE
            except ValueError:
                continue
        int_count = 0
        float_count = 0
        for s in samples:
            cleaned = s.replace(",", "")
            try:
                int(cleaned)
                int_count += 1
                continue
            except ValueError:
                pass
            try:
                float(cleaned)
                float_count += 1
            except ValueError:
                pass
        total = len(samples)
        if int_count / max(total, 1) > 0.8:
            return DataType.INTEGER
        if (int_count + float_count) / max(total, 1) > 0.7:
            return DataType.FLOAT
        return DataType.STRING

    def _detect_format_pattern(self, values: List[str], data_type: DataType) -> Optional[str]:
        if not values:
            return None
        if data_type == DataType.DATE:
            fmts = ["%Y-%m-%d", "%m/%d/%Y", "%d/%m/%Y"]
            for fmt in fmts:
                try:
                    time.strptime(values[0], fmt)
                    return fmt.replace("%Y", "YYYY").replace("%m", "MM").replace("%d", "DD")
                except ValueError:
                    continue
        if data_type == DataType.DATETIME:
            return "YYYY-MM-DDTHH:MM:SS"
        if data_type == DataType.CURRENCY:
            symbols = set()
            for v in values[:20]:
                for ch in v:
                    if ch in "$€£¥":
                        symbols.add(ch)
            sym = symbols.pop() if symbols else "$"
            return f"{sym}#,##0.00"
        if data_type == DataType.PERCENTAGE:
            return "#,##0.0%"
        return None

    def _detect_primary_key(self, columns: List[ColumnProfile]) -> List[str]:
        candidates: List[Tuple[str, float]] = []
        for col in columns:
            if col.total_count == 0:
                continue
            uniqueness = col.unique_count / col.total_count
            name_score = 1.0 if col.column_name.lower() in ("id", "key", "pk", "uuid") else 0.0
            score = uniqueness * 0.7 + name_score * 0.3
            if uniqueness >= 0.95 and col.null_count == 0:
                candidates.append((col.column_name, score))
        candidates.sort(key=lambda x: x[1], reverse=True)
        return [c[0] for c in candidates[:3]]

    def _detect_foreign_keys(self, profiles: List[SchemaProfile]) -> List[Dict[str, str]]:
        fks: List[Dict[str, str]] = []
        for p1 in profiles:
            for p2 in profiles:
                if p1.schema_id == p2.schema_id:
                    continue
                for c1 in p1.columns:
                    for c2 in p2.columns:
                        if c1.column_name == c2.column_name and c1.unique_count > 0:
                            overlap = c1.unique_count + c2.unique_count
                            if c1.null_count == 0 and c2.null_count == 0:
                                fks.append({
                                    "from_schema": p1.schema_id,
                                    "from_column": c1.column_name,
                                    "to_schema": p2.schema_id,
                                    "to_column": c2.column_name,
                                })
        return fks

    def _compute_quality_score(self, columns: List[ColumnProfile], issues: List[DataQualityIssue]) -> float:
        if not columns:
            return 100.0
        scores: List[float] = []
        for col in columns:
            completeness = 1.0 - (col.null_count / max(col.total_count, 1))
            uniqueness = col.unique_count / max(col.total_count, 1)
            col_score = (completeness * 60 + uniqueness * 40)
            scores.append(col_score)
        base = statistics.mean(scores) if scores else 100.0
        penalty = len(issues) * 2.0
        critical = sum(1 for i in issues if i.severity == "critical")
        penalty += critical * 5.0
        return round(max(0.0, min(100.0, base - penalty)), 2)

    def _detect_quality_issues(self, columns: List[ColumnProfile], rows: int) -> List[DataQualityIssue]:
        issues: List[DataQualityIssue] = []
        for col in columns:
            if col.null_count > 0:
                pct = col.null_count / max(col.total_count, 1) * 100
                sev = "critical" if pct > 50 else "high" if pct > 20 else "medium" if pct > 5 else "low"
                issues.append(DataQualityIssue(
                    issue_id=f"mv_{col.column_name}", issue_type="missing_values",
                    column=col.column_name, row=None,
                    description=f"{col.null_count} missing values ({pct:.1f}%)",
                    severity=sev, suggestion="Consider imputation or removal",
                    auto_fixable=pct < 5,
                ))
            if col.data_type in (DataType.INTEGER, DataType.FLOAT, DataType.CURRENCY) and col.std_dev and col.std_dev > 0:
                outliers = self._detect_outliers([col.mean_value or 0])
                if outliers:
                    issues.append(DataQualityIssue(
                        issue_id=f"out_{col.column_name}", issue_type="outliers",
                        column=col.column_name, row=None,
                        description=f"Potential outliers detected in {col.column_name}",
                        severity="medium", suggestion="Review outlier values",
                        auto_fixable=False,
                    ))
            if col.unique_count < col.total_count and col.column_name.lower() in ("id", "key", "uuid", "pk"):
                issues.append(DataQualityIssue(
                    issue_id=f"dup_{col.column_name}", issue_type="duplicate_keys",
                    column=col.column_name, row=None,
                    description=f"Duplicate values in key column {col.column_name}",
                    severity="critical", suggestion="Resolve duplicates to ensure data integrity",
                    auto_fixable=False,
                ))
        return issues

    def _detect_outliers(self, values: List[float]) -> List[int]:
        if len(values) < 4:
            return []
        sorted_vals = sorted(values)
        q1 = sorted_vals[len(sorted_vals) // 4]
        q3 = sorted_vals[3 * len(sorted_vals) // 4]
        iqr = q3 - q1
        lower = q1 - 1.5 * iqr
        upper = q3 + 1.5 * iqr
        return [i for i, v in enumerate(values) if v < lower or v > upper]

    def compare_schemas(self, profile1: SchemaProfile, profile2: SchemaProfile) -> Dict:
        cols1 = {c.column_name: c for c in profile1.columns}
        cols2 = {c.column_name: c for c in profile2.columns}
        added = [c for c in cols2 if c not in cols1]
        removed = [c for c in cols1 if c not in cols2]
        common = set(cols1) & set(cols2)
        type_changes = []
        for c in common:
            if cols1[c].data_type != cols2[c].data_type:
                type_changes.append({
                    "column": c,
                    "from": cols1[c].data_type.value,
                    "to": cols2[c].data_type.value,
                })
        return {
            "schema1": profile1.schema_id,
            "schema2": profile2.schema_id,
            "added_columns": added,
            "removed_columns": removed,
            "type_changes": type_changes,
            "row_count_diff": profile2.row_count - profile1.row_count,
            "quality_score_diff": round(profile2.data_quality_score - profile1.data_quality_score, 2),
        }

    def suggest_calculations(self, profile: SchemaProfile) -> List[Dict]:
        suggestions: List[Dict] = []
        numeric_cols = [c.column_name for c in profile.columns if c.data_type in (DataType.INTEGER, DataType.FLOAT, DataType.CURRENCY)]
        for i, c1 in enumerate(numeric_cols):
            suggestions.append({"type": "sum", "column": c1, "formula": f"SUM({c1})"})
            suggestions.append({"type": "average", "column": c1, "formula": f"AVERAGE({c1})"})
            suggestions.append({"type": "count", "column": c1, "formula": f"COUNT({c1})"})
            for c2 in numeric_cols[i+1:]:
                suggestions.append({"type": "ratio", "columns": [c1, c2], "formula": f"{c1}/{c2}"})
        text_cols = [c.column_name for c in profile.columns if c.data_type == DataType.STRING]
        for tc in text_cols:
            suggestions.append({"type": "unique_count", "column": tc, "formula": f"UNIQUE({tc})"})
        return suggestions

    def get_engine_summary(self) -> Dict:
        with self._lock:
            return {
                "engine_version": self.VERSION,
                "total_profiles": len(self._profiles),
                "profile_ids": list(self._profiles.keys()),
                "avg_quality_score": round(
                    statistics.mean([p.data_quality_score for p in self._profiles.values()])
                    if self._profiles else 0.0, 2
                ),
                "total_issues": sum(len(p.quality_issues) for p in self._profiles.values()),
            }
