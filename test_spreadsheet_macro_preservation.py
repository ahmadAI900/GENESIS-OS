"""Item 4A -- sorting a macro-enabled workbook must not destroy its macros.

`SpreadsheetManager._sort_xlsx` loaded every workbook with
`load_workbook(src)`. openpyxl only preserves `xl/vbaProject.bin` when
`keep_vba=True` is passed, so saving an `.xlsm` that way strips every VBA
macro from the destination -- while the call still returned
`success=True, verified=True`.

Reproduced 2026-10-15 on a real .xlsm package: the sort reported success
and the saved file had no vbaProject part at all.
"""
from __future__ import annotations

import sys
import zipfile
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent))

from genesis_os.spreadsheet_manager import SpreadsheetManager

VBA_PART = "xl/vbaProject.bin"
MACRO_SUFFIXES = (".xlsm", ".xltm")


def _has_vba(path: Path) -> bool:
    with zipfile.ZipFile(path) as archive:
        return VBA_PART in archive.namelist()


def _write_workbook(path: Path, suffix: str, *, with_vba: bool) -> None:
    """A minimal but genuine workbook package."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["name", "score"])
    ws.append(["alpha", 3])
    ws.append(["bravo", 1])
    ws.append(["charlie", 2])
    wb.save(path)
    if with_vba:
        with zipfile.ZipFile(path, "a", zipfile.ZIP_DEFLATED) as zf:
            zf.writestr(VBA_PART, b"\x00VBA-MACRO-SENTINEL\x00")


@pytest.mark.parametrize("suffix", [".xlsm"])
def test_macro_enabled_workbook_round_trips_with_its_vba(tmp_path, suffix):
    """The VBA project must survive the sort, and be reported preserved.

    `.xlsm` is the only macro-enabled format the dispatcher accepts
    (`sort` handles .csv/.xlsx/.xlsm), so it is the only one with a
    round-trip to prove. `.xltm` is deliberately NOT parameterised here:
    it is rejected upstream with "Only CSV/XLSX/XLSM ... supported", so a
    round-trip test for it would be asserting behaviour that does not
    exist.
    """
    src = tmp_path / f"book{suffix}"
    dst = tmp_path / f"sorted{suffix}"
    _write_workbook(src, suffix, with_vba=True)
    assert _has_vba(src), "fixture must genuinely carry a VBA project"

    result = SpreadsheetManager().sort(
        str(src), "score", descending=True, output=str(dst)
    )

    assert result["success"] is True, result
    assert _has_vba(dst), (
        "the sorted workbook lost xl/vbaProject.bin -- macros were "
        "silently destroyed while the call reported success"
    )
    assert result.get("macro_preserved") is True, result


def test_the_source_workbook_is_never_modified(tmp_path):
    """Sorting writes to the destination only; the source keeps its macros."""
    src = tmp_path / "book.xlsm"
    dst = tmp_path / "sorted.xlsm"
    _write_workbook(src, ".xlsm", with_vba=True)

    SpreadsheetManager().sort(
        str(src), "score", descending=True, output=str(dst)
    )

    assert _has_vba(src), "the source workbook was modified in place"


def test_plain_xlsx_is_unaffected(tmp_path):
    """A non-macro workbook sorts normally and gains no macro machinery."""
    src = tmp_path / "book.xlsx"
    dst = tmp_path / "sorted.xlsx"
    _write_workbook(src, ".xlsx", with_vba=False)

    result = SpreadsheetManager().sort(
        str(src), "score", descending=True, output=str(dst)
    )

    assert result["success"] is True, result
    assert result["executed"] is True, result
    assert result["verified"] is True, result
    assert dst.is_file()
    assert not _has_vba(dst)


def test_the_sort_actually_happened(tmp_path):
    """Guard against a 'preserved' result that silently did nothing."""
    import openpyxl

    src = tmp_path / "book.xlsm"
    dst = tmp_path / "sorted.xlsm"
    _write_workbook(src, ".xlsm", with_vba=True)

    result = SpreadsheetManager().sort(
        str(src), "score", descending=True, output=str(dst)
    )
    assert result["success"] is True, result

    rows = list(
        openpyxl.load_workbook(dst, keep_vba=True).active.iter_rows(
            min_row=2, values_only=True
        )
    )
    assert [r[0] for r in rows] == ["alpha", "charlie", "bravo"], rows


def test_keeping_vba_is_actually_requested(monkeypatch, tmp_path):
    """Structural: keep_vba must be driven by the source format.

    Asserts the call itself rather than only the outcome, so the flag
    cannot be dropped again while the vbaPart check happens to pass.
    """
    import openpyxl

    seen = {}
    real_load = openpyxl.load_workbook

    def spy(*args, **kwargs):
        seen["keep_vba"] = kwargs.get("keep_vba")
        return real_load(*args, **kwargs)

    monkeypatch.setattr(openpyxl, "load_workbook", spy)

    src = tmp_path / "book.xlsm"
    dst = tmp_path / "sorted.xlsm"
    _write_workbook(src, ".xlsm", with_vba=True)
    SpreadsheetManager().sort(str(src), "score", output=str(dst))
    assert seen.get("keep_vba") is True, (
        f"load_workbook was called with keep_vba={seen.get('keep_vba')!r} "
        "for a macro-enabled source"
    )

    seen.clear()
    src2 = tmp_path / "plain.xlsx"
    dst2 = tmp_path / "sorted2.xlsx"
    _write_workbook(src2, ".xlsx", with_vba=False)
    SpreadsheetManager().sort(str(src2), "score", output=str(dst2))
    assert seen.get("keep_vba") is False, (
        "a plain .xlsx must not be loaded with keep_vba=True"
    )