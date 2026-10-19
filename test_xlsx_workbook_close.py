"""Item 17D (C21): _extract_xlsx never closed the workbook.

`load_workbook(read_only=True, ...)` keeps the .zip archive handle
open until `wb.close()` (verified on openpyxl 3.1.5: after close(),
`wb._archive` prints `<zipfile.ZipFile [closed]>`). The old
`_extract_xlsx` returned -- on success AND on error -- without ever
closing the workbook: one leaked file handle per extracted xlsx, and
on Windows a handle that pins the file open so it cannot be deleted
or rewritten while the process lives.

The fix wraps the read+return in `try/finally` with a best-effort
`workbook.close()`. Every path is asserted here by injecting a fake
loader (monkeypatching `openpyxl.load_workbook`, the same pattern
test_spreadsheet_macro_preservation.py uses for keep_vba):

  - the success path closes,
  - a mid-read failure (the error-dict path) closes,
  - a close() that itself fails must NOT discard an extraction that
    already succeeded -- close is best-effort and makes no claim
    beyond "attempted",
  - a loader that raises has nothing to close (boundary),
  - a REAL workbook round-trip: close is spied and the underlying
    archive is verified closed, i.e. the handle really released.
"""
from __future__ import annotations

import ast
import pathlib
from pathlib import Path

import pytest

from genesis_os.filesystem.file_manager import FileManager

MODULE = pathlib.Path("genesis_os") / "filesystem" / "file_manager.py"


class _FakeWorksheet:
    def __init__(self, title, rows, fail_on_read=False):
        self.title = title
        self._rows = rows
        self._fail = fail_on_read

    def iter_rows(self, values_only=True):
        if self._fail:
            raise RuntimeError("read error inside the sheet")
        return iter(self._rows)


class _FakeWorkbook:
    def __init__(self, sheets, close_error=None):
        self.worksheets = sheets
        self.closed = False
        self.close_calls = 0
        self.close_error = close_error

    def close(self):
        self.close_calls += 1
        self.closed = True
        if self.close_error is not None:
            raise self.close_error


def _patch_loader(monkeypatch, workbook):
    """Route openpyxl.load_workbook to a fake (or a raiser if None)."""
    import openpyxl

    if workbook is None:
        def _raise(*args, **kwargs):
            raise FileNotFoundError("simulated unreadable xlsx")
        monkeypatch.setattr(openpyxl, "load_workbook", _raise)
    else:
        monkeypatch.setattr(
            openpyxl, "load_workbook", lambda *a, **k: workbook
        )


# --------------------------------------------------------------------------
# Every path must close
# --------------------------------------------------------------------------

def test_the_success_path_closes_the_workbook(monkeypatch, tmp_path):
    wb = _FakeWorkbook([
        _FakeWorksheet("Sheet1", [[1, 2], [3, 4]]),
        _FakeWorksheet("Sheet2", [["x"]]),
    ])
    _patch_loader(monkeypatch, wb)

    result = FileManager()._extract_xlsx(tmp_path / "book.xlsx")

    assert result["success"] is True, result
    assert result["sheets"] == {"Sheet1": [[1, 2], [3, 4]], "Sheet2": [["x"]]}
    assert wb.closed, (
        "the success path returned without closing the workbook; the "
        "read_only archive handle leaks on every extraction"
    )
    assert wb.close_calls == 1


def test_a_read_failure_still_closes_the_workbook(monkeypatch, tmp_path):
    wb = _FakeWorkbook(
        [_FakeWorksheet("Broken", [], fail_on_read=True)]
    )
    _patch_loader(monkeypatch, wb)

    result = FileManager()._extract_xlsx(tmp_path / "book.xlsx")

    assert result["success"] is False, result
    assert "read error inside the sheet" in result["message"]
    assert wb.closed, (
        "the error path returned without closing the workbook; the "
        "failure dict swallowed the read error and leaked the handle"
    )


def test_a_failing_close_does_not_discard_a_successful_extraction(
    monkeypatch, tmp_path
):
    """close() is best-effort: it was ATTEMPTED (that is the whole
    fix), but an exception from it must not turn an extraction that
    already succeeded into a failure dict -- and we never claim
    close() itself succeeded.
    """
    wb = _FakeWorkbook(
        [_FakeWorksheet("S", [["ok"]])],
        close_error=OSError("close blew up"),
    )
    _patch_loader(monkeypatch, wb)

    result = FileManager()._extract_xlsx(tmp_path / "book.xlsx")

    assert result["success"] is True, (
        f"a failing close() discarded the result: {result}"
    )
    assert result["sheets"] == {"S": [["ok"]]}
    assert wb.close_calls == 1, (
        "close() was not even attempted when it was needed most"
    )


def test_when_the_loader_itself_fails_there_is_nothing_to_close(
    monkeypatch, tmp_path
):
    """Boundary: load_workbook raised, so no workbook exists. The
    caller must still get the honest failure dict (no half-state).
    """
    _patch_loader(monkeypatch, None)

    result = FileManager()._extract_xlsx(tmp_path / "missing.xlsx")

    assert result["success"] is False, result
    assert "simulated unreadable xlsx" in result["message"]


def test_a_real_workbook_round_trip_closes_and_releases_the_handle(
    monkeypatch, tmp_path
):
    """Real openpyxl end-to-end: the workbook really is closed and
    the underlying archive really is released -- not just a fake
    saying it was.
    """
    from openpyxl import Workbook
    import openpyxl

    src = tmp_path / "real.xlsx"
    book = Workbook()
    book.active["A1"] = "alpha"
    book.active["B2"] = 2
    book.save(src)
    book.close()

    real_load = openpyxl.load_workbook
    seen = {}

    def spy(*args, **kwargs):
        wb = real_load(*args, **kwargs)
        seen["wb"] = wb
        seen["closed"] = False
        real_close = wb.close

        def tracking_close():
            seen["closed"] = True
            real_close()

        wb.close = tracking_close
        return wb

    monkeypatch.setattr(openpyxl, "load_workbook", spy)

    result = FileManager()._extract_xlsx(src)

    assert result["success"] is True, result
    assert result["sheets"]["Sheet"] == [["alpha", None], [None, 2]]
    assert seen.get("closed") is True, (
        "close() was never called on the real workbook; the archive "
        "handle stays open for the life of the process"
    )
    archive = getattr(seen["wb"], "_archive", None)
    assert archive is None or archive.fp is None, (
        "the workbook object was told to close but the .zip archive "
        "is still open (ZipFile.fp survives) -- the file handle is "
        "not actually released"
    )


# --------------------------------------------------------------------------
# Structural: close on every path, inside the error try
# --------------------------------------------------------------------------

def _extract_xlsx_node() -> ast.FunctionDef:
    tree = ast.parse(MODULE.read_text(encoding="utf-8-sig"))
    for cls in ast.walk(tree):
        if isinstance(cls, ast.ClassDef) and cls.name == "FileManager":
            for node in cls.body:
                if isinstance(node, ast.FunctionDef) and node.name == "_extract_xlsx":
                    return node
    raise AssertionError("FileManager._extract_xlsx not found")


def _contains(root: ast.AST, target: ast.AST) -> bool:
    return any(n is target for n in ast.walk(root))


def test_the_close_lives_in_a_finally_inside_the_error_path():
    node = _extract_xlsx_node()

    finalizers = [
        n for n in ast.walk(node)
        if isinstance(n, ast.Try) and n.finalbody
    ]
    assert len(finalizers) == 1, (
        f"expected exactly one try/finally in _extract_xlsx, got "
        f"{len(finalizers)} -- close() must run on every path"
    )
    finally_try = finalizers[0]
    finally_src = ast.unparse(finally_try)
    assert "workbook.close()" in finally_src, (
        "the finally block does not close the workbook"
    )

    holders = [
        n for n in ast.walk(node)
        if isinstance(n, ast.Try)
        and n.handlers
        and _contains(n, finally_try)
    ]
    assert holders, (
        "the try/finally is not inside the try/except that builds the "
        "failure dict; a read error would skip the close"
    )
    holder_src = ast.unparse(holders[0])
    assert "file.extract_xlsx" in holder_src and "False" in holder_src, (
        "the enclosing handler is not the extract_xlsx failure path"
    )


def test_the_source_records_the_item_it_fixes():
    """Raw text, because ast.unparse drops comments and the Item
    marker lives in the finally's comment."""
    node = _extract_xlsx_node()
    raw = MODULE.read_text(encoding="utf-8-sig").splitlines()
    body = "\n".join(raw[node.lineno - 1: node.end_lineno])
    assert "Item 17D (C21)" in body, (
        "the finally block does not record Item 17D (C21), so nobody "
        "can tell from the source why the close is mandatory"
    )
