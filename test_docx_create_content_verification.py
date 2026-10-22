"""Item 23D (C29): create_docx "verification" was an existence check.

`DocumentManager.create_docx` ended with:

    "verified": target.is_file()

-- one `os.path`-style question standing in for "the document we were
asked to create is on disk with the content we were asked to write".
`is_file()` is True for an empty file, a file a concurrent process
truncated, a corrupt non-zip blob, and a docx whose paragraphs say
something completely different. All of those claimed verified=True.

The contract asserted here:

    real python-docx write with the requested content
        -> verified=True   (reopened, paragraphs read back, matched)
    empty / wrong-content / corrupt file after save
        -> verified=False  (existence is not content)
    any exception while reopening (python-docx raises
    PackageNotFoundError on empty or corrupt packages)
        -> verified=False  (fail closed)

`verified` now also carries the evidence the executor's EvidenceContract
requires (`content_match` / `file_exists`, the keys
has_meaningful_result_evidence accepts for document.* writes), plus the
Item 13D `trustworthy_success` and `verification_type` fields -- without
the evidence key, a legitimately content-verified write was demoted to
`verification_unavailable` by _enforce_execution_contract (recorded as
the pre-fix baseline: "verified result lacks capability-specific
evidence").

TRAP pinned by the fake-docx tests: the saboteur's `Document(path)`
delegates to the REAL python-docx reader, so reading an empty or
corrupt file raises exactly the way it does in production -- the fake
sabotages the WRITE, never the readback.
"""
from __future__ import annotations

import sys
import types
from pathlib import Path

import docx as real_docx

from genesis_os.document_manager import DocumentManager


SECTIONS = [
    {"heading": "Introduction", "body": "First body paragraph."},
    {"heading": "", "body": "Body with no heading."},
    {"heading": "Heading only", "body": ""},
    "A plain-string section body.",
]


class _FakeBuilder:
    """Stands in for docx.Document() used to BUILD a document.

    Mirrors the write API create_docx uses: add_heading / add_paragraph
    record what was requested, save() then sabotages the file per mode.
    """

    def __init__(self, mode, real_module):
        self._mode = mode
        self._real = real_module
        self._recorded = []

    def add_heading(self, text, level=0):
        self._recorded.append(("heading", str(text), level))

    def add_paragraph(self, text):
        self._recorded.append(("paragraph", str(text), None))

    def save(self, target):
        if self._mode == "empty":
            Path(target).write_bytes(b"")
        elif self._mode == "corrupt":
            Path(target).write_bytes(b"this is not a zip package at all")
        elif self._mode == "raise":
            raise RuntimeError("simulated disk-full during save")
        elif self._mode == "wrong":
            doc = self._real.Document()
            doc.add_heading("A TOTALLY DIFFERENT TITLE", level=0)
            doc.add_paragraph("A TOTALLY DIFFERENT BODY")
            doc.save(str(target))
        else:  # faithful: persist the recorded content for real
            doc = self._real.Document()
            for kind, text, level in self._recorded:
                if kind == "heading":
                    doc.add_heading(text, level=level)
                else:
                    doc.add_paragraph(text)
            doc.save(str(target))


def _fake_docx_module(mode):
    """A stand-in `docx` module whose Document() branches on call shape.

    Document() with no path -> builder (the WRITE).
    Document(path)          -> the REAL python-docx reader, so opening
                               an empty or corrupt file raises exactly
                               as it does in production.
    """
    def Document(path=None):
        if path is None:
            return _FakeBuilder(mode, real_docx)
        return real_docx.Document(str(path))

    return types.SimpleNamespace(Document=Document)


def _manager_with_fake_docx(monkeypatch, mode):
    monkeypatch.setitem(sys.modules, "docx", _fake_docx_module(mode))
    return DocumentManager()


# --------------------------------------------------------------------------
# The guard: a real write, really read back.
# --------------------------------------------------------------------------

def test_a_real_docx_write_is_verified_by_content_readback(tmp_path):
    manager = DocumentManager()
    target = tmp_path / "real_write.docx"

    result = manager.create_docx(str(target), "Quarterly Report", SECTIONS)

    assert result["success"] is True, result
    assert result["executed"] is True, result
    assert result["verified"] is True, (
        "the file was written with exactly the requested content and "
        f"read back: {result}"
    )
    assert result["file_exists"] is True, result

    # Belt: the file on disk really carries the content we asked for.
    doc = real_docx.Document(str(target))
    texts = [p.text for p in doc.paragraphs]
    for expected in ("Quarterly Report", "Introduction",
                     "First body paragraph."):
        assert expected in texts, f"{expected!r} missing from {texts!r}"


def test_the_readback_verifies_content_not_existence(tmp_path):
    """The existence check alone must not be what verified means now."""
    manager = DocumentManager()
    target = tmp_path / "type_check.docx"

    result = manager.create_docx(str(target), "Typed Title", SECTIONS)

    assert result["verified"] is True, result
    assert result["verification_type"] == "docx_content_readback", (
        f"verification_type missing or wrong: {result.get('verification_type')!r}"
    )
    assert "trustworthy_success" in result, (
        f"the Item 13D field is missing; keys={sorted(result)}"
    )
    assert result["trustworthy_success"] is True, result
    assert result["content_match"] is True, (
        "the evidence the EvidenceContract accepts for a content-"
        f"verified write is missing; keys={sorted(result)}"
    )


def test_a_faithfully_faked_docx_write_still_verifies(tmp_path, monkeypatch):
    """Guard for the fake's reader delegation: honest content, faked
    module. The WRITE goes through the fake builder, the READBACK goes
    through the real python-docx reader -- verification must succeed
    exactly as with the real module, proving the helper reads via
    `from docx import Document` at call time.
    """
    manager = _manager_with_fake_docx(monkeypatch, "faithful")
    target = tmp_path / "faithful_fake.docx"

    result = manager.create_docx(str(target), "Faithful Title",
                                 [{"heading": "Section",
                                   "body": "Faithful body."}])

    assert result["success"] is True, result
    assert result["executed"] is True, result
    assert result["verified"] is True, (
        "the faked write persisted exactly the requested content and "
        f"the real reader confirmed it: {result}"
    )


# --------------------------------------------------------------------------
# The defect: every one of these claimed verified=True via is_file().
# --------------------------------------------------------------------------

def test_an_empty_written_file_is_not_verified(tmp_path, monkeypatch):
    """save() produced a zero-byte file: exists, holds nothing."""
    manager = _manager_with_fake_docx(monkeypatch, "empty")
    target = tmp_path / "empty_file.docx"

    result = manager.create_docx(str(target), "Title", SECTIONS)

    assert result["executed"] is True, (
        "save() ran without exception; the write was executed"
    )
    assert result["verified"] is False, (
        f"verified={result['verified']!r} for a zero-byte file -- "
        "existence is not content"
    )
    assert result["trustworthy_success"] is False, result
    assert result["file_exists"] is True, (
        "the file does exist; it is the CONTENT that failed, and "
        "file_exists must say so honestly"
    )


def test_a_docx_with_wrong_content_is_not_verified(tmp_path, monkeypatch):
    """save() wrote a valid docx -- with completely different content."""
    manager = _manager_with_fake_docx(monkeypatch, "wrong")
    target = tmp_path / "wrong_content.docx"

    result = manager.create_docx(str(target), "Requested Title",
                                 [{"heading": "Requested heading",
                                   "body": "Requested body."}])

    assert result["executed"] is True, result
    assert result["verified"] is False, (
        f"verified={result['verified']!r} for a docx whose paragraphs "
        "say something else entirely -- is_file() cannot see content"
    )
    assert result["content_match"] is False, result


def test_a_corrupt_written_file_is_not_verified(tmp_path, monkeypatch):
    """save() wrote bytes that are not a docx package at all.

    python-docx raises PackageNotFoundError when asked to open this;
    verification must fail closed on that exception, not crash and not
    credit the write.
    """
    manager = _manager_with_fake_docx(monkeypatch, "corrupt")
    target = tmp_path / "corrupt_file.docx"

    result = manager.create_docx(str(target), "Title", SECTIONS)

    assert result["success"] is False, (
        "save() returned without exception but verification failed, "
        "so success must reflect verified=False"
    )
    assert result["executed"] is True, result
    assert result["verified"] is False, (
        f"verified={result['verified']!r} for a file python-docx "
        "cannot even open"
    )
    assert result["trustworthy_success"] is False, result


# --------------------------------------------------------------------------
# The untouched path: a save exception is still a failure.
# --------------------------------------------------------------------------

def test_a_save_exception_still_reports_a_plain_failure(tmp_path, monkeypatch):
    manager = _manager_with_fake_docx(monkeypatch, "raise")
    target = tmp_path / "never_written.docx"

    result = manager.create_docx(str(target), "Title", SECTIONS)

    assert result["success"] is False, result
    assert result["executed"] is False, (
        "save() raised; the write was never carried out"
    )
    assert result["verified"] is False, result
    assert "simulated disk-full" in result["error"]
    assert not target.exists()
